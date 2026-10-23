import os
import re
import shutil
import subprocess
import mimetypes
import secrets
from datetime import date, datetime, time
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from dotenv import load_dotenv

load_dotenv(Path(__file__).resolve().parent / ".env")

from dateutil.relativedelta import relativedelta
from flask import (
    Flask,
    abort,
    flash,
    g,
    jsonify,
    redirect,
    render_template,
    request,
    send_from_directory,
    session,
    url_for,
    send_file,
)
from functools import wraps
from werkzeug.local import LocalProxy
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import or_, text as sql_text
from openpyxl import Workbook

from io import BytesIO
from PIL import Image
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas
from loan_rules import (
    PROTECTED_LOAN_RULES_VERSION,
    build_installment_schedule as protected_build_installment_schedule,
    calculate_multi_installment_amount as protected_calculate_multi_installment_amount,
    compute_loan_snapshot as protected_compute_loan_snapshot,
    get_next_due_date_after_interest_renegotiation,
)
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename
from storage import (
    storage_backend,
    storage_delete,
    storage_exists,
    storage_get,
    storage_key,
    storage_save,
    storage_write_bytes,
)

BASE_DIR = Path(__file__).resolve().parent
IS_VERCEL = os.getenv('VERCEL') == '1'
DATA_DIR = Path(os.getenv('LOCAL_DATA_DIR', BASE_DIR)).expanduser().resolve()
LEGACY_DB_PATH = DATA_DIR / 'mdcred.db'
DB_PATH = DATA_DIR / 'controlcred.db'
if not DB_PATH.exists() and LEGACY_DB_PATH.exists():
    shutil.copy2(LEGACY_DB_PATH, DB_PATH)
UPLOAD_DIR = DATA_DIR / 'uploads'
BACKUP_DB_DIR = DATA_DIR / 'backups' / 'db'
BACKUP_EXCEL_DIR = DATA_DIR / 'backups' / 'excel'
PUBLIC_STATIC_DIR = BASE_DIR / 'public' / 'static'
_HOSTED_PASSWORDS_CHECKED = False
ALLOWED_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.webp', '.pdf'}
LOGO_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.webp', '.svg'}
BRAND_LOGO_BASENAME = 'custom-logo'
RESERVED_USERNAMES = {'admin', 'superadmin', 'administrator', 'root'}

DOCUMENT_FIELDS = {
    'rg_foto': 'Foto do RG',
    'selfie_documento': 'Selfie com documento',
    'comprovante_residencia': 'Comprovante de residência',
    'foto_residencia': 'Foto da residência',
    'carteira_registro': 'Carteira de registro',
    'comprovante_renda': 'Comprovante de renda',
    'comprovante_pix': 'Comprovante PIX enviado',
}

app = Flask(__name__, static_folder=None if IS_VERCEL else 'static')
secret_key = os.getenv('SECRET_KEY')
if IS_VERCEL and not secret_key:
    raise RuntimeError('Configure SECRET_KEY nas variáveis de ambiente do Vercel.')
app.config['SECRET_KEY'] = secret_key or 'local-only-controlcred-secret-key'
database_url = os.getenv('DATABASE_URL')
if database_url:
    parsed_database_url = urlsplit(database_url)
    query = [(key, value) for key, value in parse_qsl(parsed_database_url.query, keep_blank_values=True)
             if key.lower() != 'pgbouncer']
    database_url = urlunsplit(parsed_database_url._replace(query=urlencode(query)))
    if database_url.startswith('postgres://'):
        database_url = database_url.replace('postgres://', 'postgresql+psycopg://', 1)
    elif database_url.startswith('postgresql://'):
        database_url = database_url.replace('postgresql://', 'postgresql+psycopg://', 1)
elif IS_VERCEL:
    raise RuntimeError('Configure DATABASE_URL com um PostgreSQL persistente antes de publicar.')
else:
    database_url = f'sqlite:///{DB_PATH}'
app.config['SQLALCHEMY_DATABASE_URI'] = database_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['UPLOAD_FOLDER'] = str(UPLOAD_DIR)
app.config['PROTECTED_LOAN_RULES_VERSION'] = '2026-03-28-rules-freeze-v1'
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['SESSION_COOKIE_SECURE'] = IS_VERCEL
app.config['MAX_CONTENT_LENGTH'] = (4 if IS_VERCEL else 10) * 1024 * 1024

if IS_VERCEL and storage_backend() != 'supabase':
    raise RuntimeError('Configure STORAGE_BACKEND=supabase e o bucket privado antes de publicar.')
if IS_VERCEL:
    missing_storage_settings = [
        name for name in ('SUPABASE_URL', 'SUPABASE_SERVICE_ROLE_KEY', 'SUPABASE_STORAGE_BUCKET')
        if not os.getenv(name)
    ]
    if missing_storage_settings:
        raise RuntimeError('Configure as variáveis SUPABASE_* antes de publicar.')

if IS_VERCEL:
    @app.route('/static/<path:filename>', endpoint='static')
    def vercel_static_file(filename):
        return send_from_directory(PUBLIC_STATIC_DIR, filename)

engine_options = {'pool_pre_ping': True}
if IS_VERCEL and database_url.startswith('postgresql+'):
    engine_options.update({'pool_size': 2, 'max_overflow': 0, 'pool_recycle': 300})
db = SQLAlchemy(app, engine_options=engine_options)


def csrf_token():
    token = session.get('_csrf_token')
    if not token:
        token = secrets.token_urlsafe(32)
        session['_csrf_token'] = token
    return token


@app.context_processor
def inject_csrf_token():
    return {'csrf_token': csrf_token}


@app.before_request
def protect_state_changes():
    if request.method in {'POST', 'PUT', 'PATCH', 'DELETE'}:
        expected = session.get('_csrf_token', '')
        supplied = request.form.get('_csrf_token', '') or request.headers.get('X-CSRF-Token', '')
        if not expected or not supplied or not secrets.compare_digest(expected, supplied):
            abort(400)


class AnonymousUser:
    id = None
    username = ''
    full_name = ''
    role = 'anon'
    charge_message_template = ''
    is_active_user = False
    is_authenticated = False

    def is_admin(self):
        return False

    def is_super_admin(self):
        return False

    def is_collector(self):
        return False

    def display_name(self):
        return ''

    def check_password(self, _password):
        return False

    def set_password(self, _password):
        return None


def _load_current_user():
    user_id = session.get('user_id')
    if not user_id:
        return AnonymousUser()
    if getattr(g, '_current_user', None) is not None:
        return g._current_user
    user = db.session.get(User, int(user_id))
    if not user or not user.is_active_user:
        session.pop('user_id', None)
        g._current_user = AnonymousUser()
        return g._current_user
    g._current_user = user
    return user


current_user = LocalProxy(_load_current_user)


class User(db.Model):
    __tablename__ = 'user'

    id = db.Column(db.Integer, primary_key=True)
    tenant_id = db.Column(db.Integer, db.ForeignKey('tenant.id'))
    full_name = db.Column(db.String(120))
    username = db.Column(db.String(80), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(20), nullable=False, default='funcionario')
    is_active_user = db.Column(db.Boolean, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    charge_message_template = db.Column(db.Text)

    clients = db.relationship('Client', back_populates='owner', lazy=True)
    loans = db.relationship('Loan', back_populates='owner_user', lazy=True)
    payments = db.relationship('Payment', back_populates='owner_payment_user', lazy=True)
    cash_movements = db.relationship('CashMovement', back_populates='owner', lazy=True)
    expenses = db.relationship('Expense', back_populates='owner', lazy=True)

    @property
    def is_authenticated(self):
        return True

    def set_password(self, password):
        self.password_hash = generate_password_hash(password)

    def check_password(self, password):
        return check_password_hash(self.password_hash, password)

    def is_admin(self):
        return (self.role or '').lower() == 'admin'

    def is_super_admin(self):
        return (self.role or '').lower() == 'superadmin'

    def is_collector(self):
        return (self.role or '').lower() == 'cobrador'

    def display_name(self):
        return (self.full_name or '').strip() or self.username


class CashBox(db.Model):
    __tablename__ = 'cash_box'

    id = db.Column(db.Integer, primary_key=True)
    current_balance = db.Column(db.Numeric(12, 2), nullable=False, default=Decimal('0.00'))
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)



class Tenant(db.Model):
    __tablename__ = 'tenant'

    id = db.Column(db.Integer, primary_key=True)
    company_name = db.Column(db.String(150), nullable=False)
    slug = db.Column(db.String(80), unique=True, nullable=False)
    logo_path = db.Column(db.String(255))
    favicon_path = db.Column(db.String(255))
    subscription_status = db.Column(db.String(30), nullable=False, default='trial')
    trial_ends_on = db.Column(db.Date)
    subscription_ends_on = db.Column(db.Date)
    is_courtesy = db.Column(db.Boolean, default=False)
    is_active = db.Column(db.Boolean, default=True)
    notes = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

    def display_name(self):
        return (self.company_name or '').strip() or f'Tenant #{self.id}'

    def subscription_label(self):
        labels = {'trial': 'Teste grátis', 'active': 'Ativo', 'expired': 'Vencido', 'blocked': 'Bloqueado', 'courtesy': 'Cortesia'}
        return labels.get(self.subscription_status, (self.subscription_status or '-').title())

    def access_until(self):
        return self.subscription_ends_on or self.trial_ends_on

    def has_active_access(self):
        if not self.is_active:
            return False
        limit_date = self.access_until()
        if self.subscription_status in {'active', 'courtesy'}:
            return True if not limit_date else limit_date >= date.today()
        if self.subscription_status == 'trial':
            return True if not limit_date else limit_date >= date.today()
        return False


class SharedRiskRecord(db.Model):
    __tablename__ = 'shared_risk_record'

    id = db.Column(db.Integer, primary_key=True)
    cpf = db.Column(db.String(14), index=True, nullable=False)
    full_name = db.Column(db.String(150), nullable=False)
    source_tenant_id = db.Column(db.Integer, db.ForeignKey('tenant.id'), nullable=False)
    active_loans = db.Column(db.Integer, nullable=False, default=0)
    active_amount = db.Column(db.Numeric(12, 2), nullable=False, default=Decimal('0.00'))
    has_overdue = db.Column(db.Boolean, default=False)
    last_updated_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)

    source_tenant = db.relationship('Tenant')

class Client(db.Model):
    __tablename__ = 'client'

    id = db.Column(db.Integer, primary_key=True)
    tenant_id = db.Column(db.Integer, db.ForeignKey('tenant.id'))
    owner_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    full_name = db.Column(db.String(150), nullable=False)
    cpf = db.Column(db.String(14), index=True, nullable=False)
    street = db.Column(db.String(150), nullable=False)
    street_number = db.Column(db.String(20))
    neighborhood = db.Column(db.String(100), nullable=False)
    city = db.Column(db.String(100), nullable=False)
    cep = db.Column(db.String(20), nullable=False)
    whatsapp = db.Column(db.String(20), nullable=False)
    phone = db.Column(db.String(20))
    contact_name = db.Column(db.String(120))
    kin_phone = db.Column(db.String(20))
    relationship = db.Column(db.String(80))
    family_street = db.Column(db.String(150))
    family_number = db.Column(db.String(20))
    family_neighborhood = db.Column(db.String(100))
    family_city = db.Column(db.String(100))
    family_cep = db.Column(db.String(20))
    family_address = db.Column(db.String(200))
    house_status = db.Column(db.String(50))
    facebook = db.Column(db.String(120))
    instagram = db.Column(db.String(120))
    email = db.Column(db.String(120))
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    rg_foto = db.Column(db.String(255))
    selfie_documento = db.Column(db.String(255))
    comprovante_residencia = db.Column(db.String(255))
    foto_residencia = db.Column(db.String(255))
    carteira_registro = db.Column(db.String(255))
    comprovante_renda = db.Column(db.String(255))
    comprovante_pix = db.Column(db.String(255))

    owner = db.relationship('User', back_populates='clients')
    loans = db.relationship('Loan', back_populates='client', cascade='all, delete-orphan', lazy=True)

    def full_address(self):
        parts = [self.street, self.street_number, self.neighborhood, self.city, self.cep]
        parts = [str(p).strip() for p in parts if p and str(p).strip()]
        return ', '.join(parts) if parts else '-'

    def family_address_display(self):
        if self.family_address and self.family_address.strip():
            return self.family_address.strip()
        parts = [self.family_street, self.family_number, self.family_neighborhood, self.family_city, self.family_cep]
        parts = [str(p).strip() for p in parts if p and str(p).strip()]
        return ', '.join(parts)

    def whatsapp_link(self, message=''):
        from urllib.parse import quote
        number = digits_only(self.whatsapp or '')
        text = quote(message or '')
        return f'https://wa.me/55{number[-11:]}?text={text}' if number else 'https://wa.me/'


class Loan(db.Model):
    __tablename__ = 'loan'

    id = db.Column(db.Integer, primary_key=True)
    tenant_id = db.Column(db.Integer, db.ForeignKey('tenant.id'))
    client_id = db.Column(db.Integer, db.ForeignKey('client.id'), nullable=False)
    owner_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    amount = db.Column(db.Numeric(12, 2), nullable=False)
    interest_percent = db.Column(db.Numeric(5, 2), nullable=False)
    loan_type = db.Column(db.String(20), nullable=False, default='mensal')
    term_days = db.Column(db.Integer)
    term_weeks = db.Column(db.Integer)
    late_fee_amount = db.Column(db.Numeric(12, 2), nullable=False, default=Decimal('20.00'))
    late_fee_interval_days = db.Column(db.Integer, nullable=False, default=1)
    charge_on_sundays = db.Column(db.Boolean, default=True)
    loan_date = db.Column(db.Date, nullable=False)
    due_date = db.Column(db.Date, nullable=False)
    status = db.Column(db.String(30), nullable=False, default='ativo')
    observations = db.Column(db.Text)
    is_active = db.Column(db.Boolean, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

    client = db.relationship('Client', back_populates='loans')
    owner_user = db.relationship('User', back_populates='loans')
    payments = db.relationship('Payment', back_populates='loan', cascade='all, delete-orphan', lazy=True)

    def interest_amount(self):
        return quantize_money(Decimal(self.amount or 0) * (Decimal(self.interest_percent or 0) / Decimal('100')))

    def snapshot(self, ref_date=None):
        return compute_loan_snapshot(self, ref_date=ref_date)


class Payment(db.Model):
    __tablename__ = 'payment'

    id = db.Column(db.Integer, primary_key=True)
    tenant_id = db.Column(db.Integer, db.ForeignKey('tenant.id'))
    loan_id = db.Column(db.Integer, db.ForeignKey('loan.id'), nullable=False)
    owner_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    amount_paid = db.Column(db.Numeric(12, 2), nullable=False)
    payment_type = db.Column(db.String(30), nullable=False)
    payment_date = db.Column(db.Date, nullable=False)
    observation = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    prev_due_date = db.Column(db.Date)
    prev_status = db.Column(db.String(30))
    prev_is_active = db.Column(db.Boolean)
    prev_observations = db.Column(db.Text)

    loan = db.relationship('Loan', back_populates='payments')
    owner_payment_user = db.relationship('User', back_populates='payments')


class CashMovement(db.Model):
    __tablename__ = 'cash_movement'

    id = db.Column(db.Integer, primary_key=True)
    tenant_id = db.Column(db.Integer, db.ForeignKey('tenant.id'))
    movement_type = db.Column(db.String(30), nullable=False)
    amount = db.Column(db.Numeric(12, 2), nullable=False)
    description = db.Column(db.String(255), nullable=False)
    reference_type = db.Column(db.String(30))
    reference_id = db.Column(db.Integer)
    owner_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

    owner = db.relationship('User', back_populates='cash_movements')


class Expense(db.Model):
    __tablename__ = 'expense'

    id = db.Column(db.Integer, primary_key=True)
    tenant_id = db.Column(db.Integer, db.ForeignKey('tenant.id'))
    amount = db.Column(db.Numeric(12, 2), nullable=False)
    category = db.Column(db.String(100), nullable=False)
    expense_date = db.Column(db.Date, nullable=False)
    observation = db.Column(db.Text)
    owner_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

    owner = db.relationship('User', back_populates='expenses')


class CollectorVisit(db.Model):
    __tablename__ = 'collector_visit'

    id = db.Column(db.Integer, primary_key=True)
    tenant_id = db.Column(db.Integer, db.ForeignKey('tenant.id'))
    client_id = db.Column(db.Integer, db.ForeignKey('client.id'), nullable=False)
    collector_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    status = db.Column(db.String(40), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)

    client = db.relationship('Client')
    collector = db.relationship('User')


class DailyCashClosure(db.Model):
    __tablename__ = 'daily_cash_closure'

    id = db.Column(db.Integer, primary_key=True)
    tenant_id = db.Column(db.Integer, db.ForeignKey('tenant.id'))
    closure_date = db.Column(db.Date, nullable=False)
    opening_balance = db.Column(db.Numeric(12, 2), nullable=False, default=Decimal('0.00'))
    total_inflows = db.Column(db.Numeric(12, 2), nullable=False, default=Decimal('0.00'))
    total_outflows = db.Column(db.Numeric(12, 2), nullable=False, default=Decimal('0.00'))
    ending_balance = db.Column(db.Numeric(12, 2), nullable=False, default=Decimal('0.00'))
    notes = db.Column(db.Text)
    closed_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    closed_by_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)

    closed_by = db.relationship('User')


VISIT_STATUS_LABELS = {
    'sucesso': 'Visita concluída com sucesso',
    'ausente': 'Cliente não estava presente',
    'rota': 'Não foi possível ir até essa rota',
}



def login_user(user):
    session['user_id'] = user.id
    g._current_user = user



def logout_user():
    session.pop('user_id', None)
    g._current_user = AnonymousUser()



def login_required(view_func):
    @wraps(view_func)
    def wrapped(*args, **kwargs):
        if not current_user.is_authenticated:
            return redirect(url_for('login'))
        return view_func(*args, **kwargs)
    return wrapped



def get_cashbox():
    cashbox = db.session.get(CashBox, 1)
    if not cashbox:
        cashbox = CashBox(id=1, current_balance=Decimal('0.00'))
        db.session.add(cashbox)
        db.session.flush()
    return cashbox



def total_expenses_value():
    total = db.session.query(db.func.coalesce(db.func.sum(Expense.amount), 0)).scalar()
    return quantize_money(total or 0)



def apply_cash_movement(amount, movement_type, description, reference_type=None, reference_id=None):
    amount = quantize_money(amount)
    cashbox = get_cashbox()
    cashbox.current_balance = quantize_money(Decimal(cashbox.current_balance or 0) + amount)
    movement = CashMovement(
        movement_type=movement_type,
        amount=amount,
        description=description,
        reference_type=reference_type,
        reference_id=reference_id,
        owner_id=current_user.id or 1,
    )
    db.session.add(movement)
    cashbox.updated_at = datetime.utcnow()
    return movement



def reverse_cash_movements(reference_type, reference_id):
    movements = CashMovement.query.filter_by(reference_type=reference_type, reference_id=reference_id).all()
    if not movements:
        return
    cashbox = get_cashbox()
    for movement in movements:
        cashbox.current_balance = quantize_money(Decimal(cashbox.current_balance or 0) - Decimal(movement.amount or 0))
        db.session.delete(movement)
    cashbox.updated_at = datetime.utcnow()



def sync_loan_cash_movement(loan):
    movement = CashMovement.query.filter_by(reference_type='loan', reference_id=loan.id).order_by(CashMovement.id.desc()).first()
    if movement:
        update_cash_movement_amount(movement, -Decimal(loan.amount or 0))
        movement.description = f'Empréstimo #{loan.id} para {loan.client.full_name}'



def now_br():
    return datetime.now().strftime('%d/%m/%Y %H:%M')


def fmt_money(value):
    value = Decimal(value or 0)
    return f'R$ {value:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def fmt_date(value):
    if not value:
        return '-'
    if isinstance(value, datetime):
        value = value.date()
    return value.strftime('%d/%m/%Y')


def digits_only(text):
    return re.sub(r'\D', '', text or '')


def format_cpf(value):
    nums = digits_only(value)
    if len(nums) != 11:
        return value or ''
    return f'{nums[:3]}.{nums[3:6]}.{nums[6:9]}-{nums[9:]}'


def format_phone(value):
    nums = digits_only(value)
    if len(nums) == 11:
        return f'({nums[:2]}) {nums[2:7]}-{nums[7:]}'
    if len(nums) == 10:
        return f'({nums[:2]}) {nums[2:6]}-{nums[6:]}'
    return value or ''


def normalize_whatsapp(value):
    nums = digits_only(value)
    if nums.startswith('55'):
        nums = nums[2:]
    if len(nums) not in (10, 11):
        return value or ''
    return '+55' + nums




def default_charge_message_template():
    return (
        'Olá {nome}, tudo bem?\n\n'
        'Identificamos um débito no valor de R$ {valor}, com vencimento em {vencimento}.\n\n'
        'Atualmente está em atraso há {dias_atraso} dia(s).\n\n'
        'Por favor, entre em contato para regularizar.\n\n'
        'Obrigado!'
    )


def render_charge_message(template, client, loan=None, snapshot=None):
    snap = snapshot or (loan.snapshot() if loan is not None else None)
    if snap:
        amount = snap.get('current_due_amount') or snap.get('updated_amount') or snap.get('remaining_amount') or Decimal('0.00')
        due_date = loan.due_date if loan is not None else None
        days_overdue = int(snap.get('days_overdue') or 0)
    else:
        amount = Decimal('0.00')
        due_date = None
        days_overdue = 0
    final_template = (template or '').strip() or default_charge_message_template()
    return final_template.format(
        nome=client.full_name,
        valor=fmt_money(amount).replace('R$ ', ''),
        vencimento=fmt_date(due_date),
        dias_atraso=days_overdue,
    )


def form_value(form_data, obj, attr, default=''):
    if form_data and attr in form_data:
        return form_data.get(attr, default)
    return getattr(obj, attr, default) if obj is not None else default


def allowed_file(filename):
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS


def allowed_logo_file(filename):
    return Path(filename).suffix.lower() in LOGO_EXTENSIONS


def current_tenant():
    if not current_user.is_authenticated:
        return None
    if current_user.is_super_admin():
        tenant_id = session.get('active_tenant_id')
        return db.session.get(Tenant, int(tenant_id)) if tenant_id else None
    return db.session.get(Tenant, int(current_user.tenant_id)) if getattr(current_user, 'tenant_id', None) else None


def tenant_display_name():
    tenant = current_tenant()
    return tenant.display_name() if tenant else 'ControlCred'



def brand_logo_url():
    tenant = current_tenant()
    if tenant and tenant.logo_path:
        key = storage_key(tenant.logo_path)
        if storage_exists(key):
            return url_for('uploaded_file', filename=key, v=tenant.id)
    for ext in ('.png', '.jpg', '.jpeg', '.webp'):
        global_logo_key = f'brand/global/logo{ext}'
        if storage_exists(global_logo_key):
            return url_for('uploaded_file', filename=global_logo_key, v=int(datetime.now().timestamp()))
    return url_for('static', filename='img/logo.png')


def brand_logo_bytes():
    tenant = current_tenant()
    if tenant and tenant.logo_path:
        key = storage_key(tenant.logo_path)
        if storage_exists(key):
            return storage_get(key)
    for ext in ('.png', '.jpg', '.jpeg', '.webp'):
        key = f'brand/global/logo{ext}'
        if storage_exists(key):
            return storage_get(key)
    fallback = BASE_DIR / 'static' / 'img' / 'logo.png'
    return fallback.read_bytes() if fallback.exists() else None


def brand_favicon_url():
    tenant = current_tenant()
    if tenant and tenant.favicon_path:
        key = storage_key(tenant.favicon_path)
        if storage_exists(key):
            return url_for('uploaded_file', filename=key, v=tenant.id)
    global_favicon_key = 'brand/global/favicon.ico'
    if storage_exists(global_favicon_key):
        return url_for('uploaded_file', filename=global_favicon_key, v=int(datetime.now().timestamp()))
    static_favicon = BASE_DIR / 'static' / 'favicon.ico'
    if static_favicon.exists():
        return url_for('static', filename='favicon.ico')
    return brand_logo_url()



def save_brand_logo(file_storage, tenant=None):
    tenant = tenant or current_tenant()
    if not file_storage or not file_storage.filename:
        return None
    ext = Path(file_storage.filename).suffix.lower()
    if ext not in {'.png', '.jpg', '.jpeg', '.webp'}:
        raise ValueError('Arquivo de logotipo inválido. Envie PNG, JPG ou WEBP.')

    def generate_favicon(source_bytes: bytes) -> bytes | None:
        try:
            img = Image.open(BytesIO(source_bytes)).convert('RGBA')
            size = max(img.size)
            canvas_img = Image.new('RGBA', (size, size), (0, 0, 0, 0))
            canvas_img.paste(img, ((size - img.size[0]) // 2, (size - img.size[1]) // 2), img)
            favicon = BytesIO()
            canvas_img.save(favicon, format='ICO', sizes=[(16,16), (32,32), (48,48), (64,64)])
            return favicon.getvalue()
        except Exception:
            return None

    file_storage.stream.seek(0)
    logo_bytes = file_storage.stream.read()
    favicon_bytes = generate_favicon(logo_bytes)

    if tenant:
        logo_key = f'tenants/{tenant.id}/logo{ext}'
        favicon_key = f'tenants/{tenant.id}/favicon.ico'
        storage_write_bytes(logo_key, logo_bytes, file_storage.mimetype or 'application/octet-stream')
        if favicon_bytes:
            storage_write_bytes(favicon_key, favicon_bytes, 'image/x-icon')
        tenant.logo_path = f'uploads/{logo_key}'
        tenant.favicon_path = f'uploads/{favicon_key}'
        return logo_key

    logo_key = 'brand/global/logo' + ext
    storage_write_bytes(logo_key, logo_bytes, file_storage.mimetype or 'application/octet-stream')
    if favicon_bytes:
        storage_write_bytes('brand/global/favicon.ico', favicon_bytes, 'image/x-icon')
    return logo_key



def month_add(dt):
    return dt + relativedelta(months=1)


def quantize_money(value):
    return Decimal(value or 0).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def installment_count_for(loan):
    loan_type = getattr(loan, 'loan_type', 'mensal') or 'mensal'
    if loan_type == 'diario':
        return max(int(getattr(loan, 'term_days', 30) or 30), 1)
    if loan_type == 'semanal':
        return max(int(getattr(loan, 'term_weeks', 4) or 4), 1)
    return 1


def daily_due_dates(start_date, installments, charge_on_sundays=True):
    due_dates = []
    cursor = start_date
    while len(due_dates) < max(int(installments or 0), 0):
        cursor = cursor + relativedelta(days=1)
        if not charge_on_sundays and cursor.weekday() == 6:
            continue
        due_dates.append(cursor)
    return due_dates


def build_installment_schedule(loan, ref_date=None):
    """REGRA CONGELADA: cronograma oficial em loan_rules.py."""
    return protected_build_installment_schedule(loan, ref_date=ref_date)


def calculate_multi_installment_amount(loan, quantity, ref_date=None):
    """REGRA CONGELADA: cálculo oficial em loan_rules.py."""
    return protected_calculate_multi_installment_amount(loan, quantity, ref_date=ref_date)


def paid_installments_count(total_paid, installment_amount, total_installments):
    installment_amount = Decimal(installment_amount or 0)
    if installment_amount <= 0:
        return 0
    total_paid = Decimal(total_paid or 0)
    count = int(total_paid // installment_amount)
    return max(0, min(count, int(total_installments or 0)))

def loan_cycle_label(loan):
    loan_type = getattr(loan, 'loan_type', 'mensal') or 'mensal'
    if loan_type == 'diario':
        return 'Diário'
    if loan_type == 'semanal':
        return 'Semanal'
    return 'Mensal'


def compute_loan_snapshot(loan, ref_date=None):
    """
    REGRA CONGELADA: snapshot oficial do contrato mora em loan_rules.py.
    Versão protegida: 2026-03-28-rules-freeze-v1
    """
    return protected_compute_loan_snapshot(loan, ref_date=ref_date, installment_count_for=installment_count_for)


def calculate_client_score(client):
    score = 100
    good_closed_loans = 0

    loans = sorted(list(client.loans or []), key=lambda loan: (loan.loan_date or date.min, loan.id or 0))
    for loan in loans:
        payments = sorted(list(loan.payments or []), key=lambda payment: (payment.payment_date or date.min, payment.id or 0))
        loan_had_problem = False

        for payment in payments:
            due_ref = payment.prev_due_date or loan.due_date or payment.payment_date
            overdue_days = max((payment.payment_date - due_ref).days, 0)
            if overdue_days > 0:
                loan_had_problem = True
                if overdue_days <= 3:
                    score -= 2
                elif overdue_days <= 10:
                    score -= 5
                else:
                    score -= 10
            else:
                score += 1

            if payment.payment_type in {'negociar_juros', 'negociar_quitacao'}:
                loan_had_problem = True
                score -= 8

        snap = loan.snapshot()
        if loan.is_active and (snap.get('days_overdue') or 0) > 0:
            loan_had_problem = True
            current_delay = snap.get('days_overdue') or 0
            if current_delay <= 3:
                score -= 2
            elif current_delay <= 10:
                score -= 5
            else:
                score -= 10

        fine_value = Decimal(snap.get('fine') or 0)
        if fine_value > 0:
            if fine_value <= Decimal('100'):
                score -= 2
            elif fine_value <= Decimal('500'):
                score -= 5
            else:
                score -= 10

        if not loan.is_active and payments and not loan_had_problem:
            good_closed_loans += 1
            score += 10

    if good_closed_loans >= 3:
        score += 15

    visits = CollectorVisit.query.filter_by(client_id=client.id).all()
    for visit in visits:
        if visit.status == 'ausente':
            score -= 3
        elif visit.status == 'rota':
            score -= 2
        elif visit.status == 'sucesso':
            score += 1

    score = max(0, min(100, int(score)))
    if score >= 80:
        label = 'Excelente'
        css_class = 'score-excellent'
    elif score >= 60:
        label = 'Bom'
        css_class = 'score-good'
    elif score >= 40:
        label = 'Risco médio'
        css_class = 'score-medium'
    else:
        label = 'Alto risco'
        css_class = 'score-high'

    return {'value': score, 'label': label, 'css_class': css_class}


@app.context_processor
def inject_helpers():
    return dict(
        fmt_money=fmt_money,
        fmt_date=fmt_date,
        loan_snapshot=compute_loan_snapshot,
        is_image=lambda path: Path(path or '').suffix.lower() in {'.png', '.jpg', '.jpeg', '.webp'},
        document_fields=DOCUMENT_FIELDS,
        getattr=getattr,
        current_user=current_user,
        form_value=form_value,
        loan_cycle_label=loan_cycle_label,
        render_charge_message=render_charge_message,
        default_charge_message_template=default_charge_message_template,
        brand_logo_url=brand_logo_url,
        brand_favicon_url=brand_favicon_url,
        tenant_display_name=tenant_display_name,
        current_tenant=current_tenant,
        today=date.today,
        visit_status_labels=VISIT_STATUS_LABELS,
        client_score=calculate_client_score,
    )


def ensure_dirs():
    if IS_VERCEL:
        return
    for folder in [UPLOAD_DIR, BACKUP_DB_DIR, BACKUP_EXCEL_DIR, UPLOAD_DIR / 'tenants']:
        folder.mkdir(parents=True, exist_ok=True)


def create_initial_data():
    db.create_all()
    principal = Tenant.query.filter_by(slug='principal').first()
    if not principal:
        principal = Tenant(company_name='ControlCred Principal', slug='principal', subscription_status='courtesy', is_courtesy=True, is_active=True, subscription_ends_on=date.today() + relativedelta(years=10))
        db.session.add(principal)
        db.session.commit()
    admin = User.query.filter_by(username='admin').first()
    if not admin:
        admin = User(username='admin', full_name='Administrador da Plataforma', role='superadmin', is_active_user=True)
        admin_password = os.getenv('INITIAL_ADMIN_PASSWORD') or ('123456' if not IS_VERCEL else '')
        if not admin_password:
            raise RuntimeError('Defina INITIAL_ADMIN_PASSWORD antes de inicializar os dados no Vercel.')
        admin.set_password(admin_password)
        db.session.add(admin)
        db.session.commit()
    elif IS_VERCEL and check_password_hash(admin.password_hash, '123456'):
        admin_password = os.getenv('INITIAL_ADMIN_PASSWORD')
        if not admin_password:
            raise RuntimeError('Troque a senha padrão do admin: defina INITIAL_ADMIN_PASSWORD.')
        admin.set_password(admin_password)
        db.session.commit()
    tenant_admin = User.query.filter_by(username='tenantadmin').first()
    if not tenant_admin:
        tenant_admin = User(username='tenantadmin', full_name='Administrador do Cliente', role='admin', is_active_user=True, tenant_id=principal.id)
        tenant_password = os.getenv('INITIAL_TENANT_ADMIN_PASSWORD') or ('123456' if not IS_VERCEL else '')
        if not tenant_password:
            raise RuntimeError('Defina INITIAL_TENANT_ADMIN_PASSWORD antes de inicializar os dados no Vercel.')
        tenant_admin.set_password(tenant_password)
        db.session.add(tenant_admin)
        db.session.commit()
    elif IS_VERCEL and check_password_hash(tenant_admin.password_hash, '123456'):
        tenant_password = os.getenv('INITIAL_TENANT_ADMIN_PASSWORD')
        if not tenant_password:
            raise RuntimeError('Troque a senha padrão do tenantadmin: defina INITIAL_TENANT_ADMIN_PASSWORD.')
        tenant_admin.set_password(tenant_password)
        db.session.commit()


def run_migrations():
    # This legacy migration block uses SQLite PRAGMA/sqlite_master. Postgres
    # schemas are created from the SQLAlchemy models and imported separately.
    if db.engine.dialect.name != 'sqlite':
        db.create_all()
        return
    with db.engine.begin() as conn:
        db.metadata.create_all(bind=db.engine)

        # user
        user_columns = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(user)"))}
        if 'tenant_id' not in user_columns:
            conn.execute(sql_text("ALTER TABLE user ADD COLUMN tenant_id INTEGER"))

        # rebuild client table once to remove UNIQUE on CPF and add tenant_id
        create_sql_row = conn.execute(sql_text("SELECT sql FROM sqlite_master WHERE type='table' AND name='client'")).fetchone()
        create_sql = create_sql_row[0] if create_sql_row else ''
        if create_sql and 'UNIQUE' in create_sql.upper():
            conn.execute(sql_text("ALTER TABLE client RENAME TO client_old_multi"))
            conn.execute(sql_text("""CREATE TABLE client (
                id INTEGER PRIMARY KEY,
                tenant_id INTEGER,
                owner_id INTEGER NOT NULL,
                full_name VARCHAR(150) NOT NULL,
                cpf VARCHAR(14) NOT NULL,
                street VARCHAR(150) NOT NULL,
                street_number VARCHAR(20),
                neighborhood VARCHAR(100) NOT NULL,
                city VARCHAR(100) NOT NULL,
                cep VARCHAR(20) NOT NULL,
                whatsapp VARCHAR(20) NOT NULL,
                phone VARCHAR(20),
                contact_name VARCHAR(120),
                kin_phone VARCHAR(20),
                relationship VARCHAR(80),
                family_street VARCHAR(150),
                family_number VARCHAR(20),
                family_neighborhood VARCHAR(100),
                family_city VARCHAR(100),
                family_cep VARCHAR(20),
                family_address VARCHAR(200),
                house_status VARCHAR(50),
                facebook VARCHAR(120),
                instagram VARCHAR(120),
                email VARCHAR(120),
                created_at DATETIME,
                rg_foto VARCHAR(255),
                selfie_documento VARCHAR(255),
                comprovante_residencia VARCHAR(255),
                foto_residencia VARCHAR(255),
                carteira_registro VARCHAR(255),
                comprovante_renda VARCHAR(255),
                comprovante_pix VARCHAR(255)
            )"""))
            conn.execute(sql_text("""INSERT INTO client (
                id, tenant_id, owner_id, full_name, cpf, street, street_number, neighborhood, city, cep, whatsapp, phone, contact_name, kin_phone, relationship, family_street, family_number, family_neighborhood, family_city, family_cep, family_address, house_status, facebook, instagram, email, created_at, rg_foto, selfie_documento, comprovante_residencia, foto_residencia, carteira_registro, comprovante_renda, comprovante_pix
            )
            SELECT c.id, u.tenant_id, c.owner_id, c.full_name, c.cpf, c.street, c.street_number, c.neighborhood, c.city, c.cep, c.whatsapp, c.phone, c.contact_name, c.kin_phone, c.relationship, c.family_street, c.family_number, c.family_neighborhood, c.family_city, c.family_cep, c.family_address, c.house_status, c.facebook, c.instagram, c.email, c.created_at, c.rg_foto, c.selfie_documento, c.comprovante_residencia, c.foto_residencia, c.carteira_registro, c.comprovante_renda, c.comprovante_pix
            FROM client_old_multi c LEFT JOIN user u ON u.id = c.owner_id"""))
            conn.execute(sql_text("DROP TABLE client_old_multi"))
            conn.execute(sql_text("CREATE INDEX IF NOT EXISTS ix_client_cpf ON client (cpf)"))

        # add tenant_id columns to remaining tables
        table_ddls = {
            ('client','tenant_id'): "ALTER TABLE client ADD COLUMN tenant_id INTEGER",
            ('loan','tenant_id'): "ALTER TABLE loan ADD COLUMN tenant_id INTEGER",
            ('payment','tenant_id'): "ALTER TABLE payment ADD COLUMN tenant_id INTEGER",
            ('cash_movement','tenant_id'): "ALTER TABLE cash_movement ADD COLUMN tenant_id INTEGER",
            ('expense','tenant_id'): "ALTER TABLE expense ADD COLUMN tenant_id INTEGER",
            ('collector_visit','tenant_id'): "ALTER TABLE collector_visit ADD COLUMN tenant_id INTEGER",
            ('daily_cash_closure','tenant_id'): "ALTER TABLE daily_cash_closure ADD COLUMN tenant_id INTEGER",
        }
        for (table_name, column), ddl in table_ddls.items():
            cols = {row[1] for row in conn.execute(sql_text(f"PRAGMA table_info({table_name})"))}
            if column not in cols:
                conn.execute(sql_text(ddl))

        # original additive migrations
        columns = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(client)"))}
        additions = {
            'street_number': 'ALTER TABLE client ADD COLUMN street_number VARCHAR(20)',
            'kin_phone': 'ALTER TABLE client ADD COLUMN kin_phone VARCHAR(20)',
            'family_street': 'ALTER TABLE client ADD COLUMN family_street VARCHAR(150)',
            'family_number': 'ALTER TABLE client ADD COLUMN family_number VARCHAR(20)',
            'family_neighborhood': 'ALTER TABLE client ADD COLUMN family_neighborhood VARCHAR(100)',
            'family_city': 'ALTER TABLE client ADD COLUMN family_city VARCHAR(100)',
            'family_cep': 'ALTER TABLE client ADD COLUMN family_cep VARCHAR(20)',
        }
        for column, ddl in additions.items():
            if column not in columns:
                conn.execute(sql_text(ddl))

        loan_columns = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(loan)"))}
        loan_additions = {
            'loan_type': "ALTER TABLE loan ADD COLUMN loan_type VARCHAR(20) DEFAULT 'mensal'",
            'term_days': 'ALTER TABLE loan ADD COLUMN term_days INTEGER DEFAULT 30',
            'term_weeks': 'ALTER TABLE loan ADD COLUMN term_weeks INTEGER DEFAULT 4',
            'late_fee_amount': 'ALTER TABLE loan ADD COLUMN late_fee_amount NUMERIC(12,2) DEFAULT 20.00',
            'late_fee_interval_days': 'ALTER TABLE loan ADD COLUMN late_fee_interval_days INTEGER DEFAULT 1',
            'charge_on_sundays': 'ALTER TABLE loan ADD COLUMN charge_on_sundays BOOLEAN DEFAULT 1',
        }
        for column, ddl in loan_additions.items():
            if column not in loan_columns:
                conn.execute(sql_text(ddl))

        user_columns = {row[1] for row in conn.execute(sql_text("PRAGMA table_info(user)"))}
        if 'charge_message_template' not in user_columns:
            conn.execute(sql_text('ALTER TABLE user ADD COLUMN charge_message_template TEXT'))
        if 'full_name' not in user_columns:
            conn.execute(sql_text('ALTER TABLE user ADD COLUMN full_name VARCHAR(120)'))

        tables = {row[0] for row in conn.execute(sql_text("SELECT name FROM sqlite_master WHERE type='table'"))}
        if 'collector_visit' not in tables:
            conn.execute(sql_text('CREATE TABLE collector_visit (id INTEGER PRIMARY KEY AUTOINCREMENT, client_id INTEGER NOT NULL, collector_id INTEGER NOT NULL, status VARCHAR(40) NOT NULL, created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP, FOREIGN KEY(client_id) REFERENCES client(id), FOREIGN KEY(collector_id) REFERENCES user(id))'))
        if 'daily_cash_closure' not in tables:
            conn.execute(sql_text("CREATE TABLE daily_cash_closure (id INTEGER PRIMARY KEY AUTOINCREMENT, closure_date DATE NOT NULL UNIQUE, opening_balance NUMERIC(12,2) NOT NULL DEFAULT 0.00, total_inflows NUMERIC(12,2) NOT NULL DEFAULT 0.00, total_outflows NUMERIC(12,2) NOT NULL DEFAULT 0.00, ending_balance NUMERIC(12,2) NOT NULL DEFAULT 0.00, notes TEXT, closed_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP, closed_by_id INTEGER NOT NULL, FOREIGN KEY(closed_by_id) REFERENCES user(id))"))


def save_upload(file_storage, client_id, field_name, old_filename=None):
    if not file_storage or not file_storage.filename:
        return old_filename
    if not allowed_file(file_storage.filename):
        raise ValueError('Arquivo inválido. Envie imagem ou PDF.')
    ext = Path(file_storage.filename).suffix.lower()
    filename = secure_filename(f'cliente_{client_id}_{field_name}_{int(datetime.now().timestamp())}{ext}')
    storage_save(file_storage, filename)
    if old_filename:
        delete_document_file(old_filename)
    return filename


def delete_document_file(filename):
    if filename:
        storage_delete(filename)


def remove_client_document(client, field_name):
    if field_name not in DOCUMENT_FIELDS:
        raise ValueError('Documento inválido.')
    current = getattr(client, field_name)
    if current:
        delete_document_file(current)
        setattr(client, field_name, None)



def delete_tenant_brand_assets(tenant):
    if not tenant:
        return
    for rel_path in [getattr(tenant, 'logo_path', None), getattr(tenant, 'favicon_path', None)]:
        if rel_path:
            storage_delete(rel_path)
    for key in (f'tenants/{tenant.id}/logo.png', f'tenants/{tenant.id}/logo.jpg', f'tenants/{tenant.id}/logo.jpeg', f'tenants/{tenant.id}/logo.webp', f'tenants/{tenant.id}/favicon.ico'):
        storage_delete(key)


def cascade_delete_tenant_data(tenant):
    if not tenant:
        return
    tenant_id = tenant.id
    clients = Client.query.filter_by(tenant_id=tenant_id).all()
    for client in clients:
        for field in DOCUMENT_FIELDS:
            delete_document_file(getattr(client, field))
    SharedRiskRecord.query.filter_by(source_tenant_id=tenant_id).delete(synchronize_session=False)
    Payment.query.filter_by(tenant_id=tenant_id).delete(synchronize_session=False)
    CollectorVisit.query.filter_by(tenant_id=tenant_id).delete(synchronize_session=False)
    DailyCashClosure.query.filter_by(tenant_id=tenant_id).delete(synchronize_session=False)
    Expense.query.filter_by(tenant_id=tenant_id).delete(synchronize_session=False)
    CashMovement.query.filter_by(tenant_id=tenant_id).delete(synchronize_session=False)
    Loan.query.filter_by(tenant_id=tenant_id).delete(synchronize_session=False)
    Client.query.filter_by(tenant_id=tenant_id).delete(synchronize_session=False)
    User.query.filter(User.tenant_id == tenant_id, User.role != 'superadmin').delete(synchronize_session=False)
    delete_tenant_brand_assets(tenant)
    db.session.delete(tenant)




@app.before_request
def protect_root():
    global _HOSTED_PASSWORDS_CHECKED
    if IS_VERCEL and request.endpoint != 'static' and not _HOSTED_PASSWORDS_CHECKED:
        for username, variable in (
            ('admin', 'INITIAL_ADMIN_PASSWORD'),
            ('tenantadmin', 'INITIAL_TENANT_ADMIN_PASSWORD'),
        ):
            user = User.query.filter_by(username=username).first()
            if user and check_password_hash(user.password_hash, '123456'):
                new_password = os.getenv(variable)
                if not new_password:
                    abort(503, description=f'Configure {variable} para substituir a senha padrão antes de acessar o sistema.')
                user.set_password(new_password)
        db.session.commit()
        _HOSTED_PASSWORDS_CHECKED = True

    public_routes = {'login', 'register', 'static', 'uploaded_file', 'desktop_installer'}
    if request.endpoint in public_routes or request.endpoint is None:
        return
    if not current_user.is_authenticated:
        return redirect(url_for('login'))
    if current_user.is_super_admin():
        return
    if not tenant_has_access() and request.endpoint not in {'logout', 'profile', 'dashboard'}:
        flash('Sua assinatura está vencida ou bloqueada. Renove para continuar usando todas as funções.', 'danger')
        return redirect(url_for('dashboard'))
    if current_user.is_collector():
        allowed = {'dashboard', 'logout', 'uploaded_file', 'collector_client_detail', 'collector_visit_update', 'profile'}
        if request.endpoint not in allowed:
            return redirect(url_for('dashboard'))


def require_admin():
    if not current_user.is_authenticated or not current_user.is_admin():
        flash('Acesso permitido apenas para administrador.', 'danger')
        return False
    return True


def require_super_admin():
    if not current_user.is_authenticated or not current_user.is_super_admin():
        flash('Acesso permitido apenas ao administrador da plataforma.', 'danger')
        return False
    return True


def active_tenant_id():
    tenant = current_tenant()
    return tenant.id if tenant else None


def filter_query_by_tenant(query, model):
    if current_user.is_super_admin():
        active = session.get('active_tenant_id')
        return query.filter(model.tenant_id == active) if active and hasattr(model, 'tenant_id') else query
    tenant_id = active_tenant_id()
    return query.filter(model.tenant_id == tenant_id) if tenant_id and hasattr(model, 'tenant_id') else query


def tenant_has_access():
    tenant = current_tenant()
    return True if current_user.is_super_admin() else bool(tenant and tenant.has_active_access())


def sync_shared_risk_for_cpf(cpf):
    cpf_digits = digits_only(cpf)
    if not cpf_digits:
        return
    SharedRiskRecord.query.filter_by(cpf=format_cpf(cpf_digits)).delete()
    clients = Client.query.all()
    grouped = {}
    for client in clients:
        if digits_only(client.cpf) != cpf_digits:
            continue
        tenant_id = getattr(client, 'tenant_id', None) or getattr(client.owner, 'tenant_id', None)
        if not tenant_id:
            continue
        data = grouped.setdefault(tenant_id, {'name': client.full_name, 'amount': Decimal('0.00'), 'count': 0, 'overdue': False})
        for loan in Loan.query.filter_by(client_id=client.id, is_active=True).all():
            data['count'] += 1
            data['amount'] += Decimal(loan.amount or 0)
            snap = loan.snapshot()
            if (snap.get('days_overdue') or 0) > 0:
                data['overdue'] = True
    for tenant_id, info in grouped.items():
        db.session.add(SharedRiskRecord(cpf=format_cpf(cpf_digits), full_name=info['name'], source_tenant_id=tenant_id, active_loans=info['count'], active_amount=quantize_money(info['amount']), has_overdue=info['overdue'], last_updated_at=datetime.utcnow()))



def risk_alert_for_cpf(cpf):
    cpf_fmt = format_cpf(cpf)
    if not cpf_fmt:
        return []
    try:
        sync_shared_risk_for_cpf(cpf_fmt)
        db.session.flush()
    except Exception:
        db.session.rollback()
    query = SharedRiskRecord.query.filter_by(cpf=cpf_fmt)
    tenant_id = active_tenant_id()
    if tenant_id:
        query = query.filter(SharedRiskRecord.source_tenant_id != tenant_id)
    return query.order_by(SharedRiskRecord.last_updated_at.desc()).all()



def build_risk_alert_payload(cpf):
    alerts = risk_alert_for_cpf(cpf)
    if not alerts:
        return {'has_alert': False, 'items': []}
    items = []
    for alert in alerts:
        active_loans = int(alert.active_loans or 0)
        active_amount = quantize_money(alert.active_amount or 0)
        if alert.has_overdue:
            status_label = 'Com atraso'
        elif active_loans > 0 or active_amount > 0:
            status_label = 'Em dia'
        else:
            status_label = 'Quitado'
        items.append({
            'active_amount': fmt_money(active_amount),
            'active_loans': active_loans,
            'status_label': status_label,
        })
    summary = items[0]
    return {
        'has_alert': True,
        'items': items,
        'message_title': 'Cliente já possui vínculo com outro credor da rede',
        'active_amount': summary['active_amount'],
        'active_loans': summary['active_loans'],
        'status_label': summary['status_label'],
    }


@app.route('/api/risk-check/client')
@login_required
def api_risk_check_client():
    cpf = request.args.get('cpf', '').strip()
    return jsonify(build_risk_alert_payload(cpf))


@app.route('/api/risk-check/loan')
@login_required
def api_risk_check_loan():
    client_id = request.args.get('client_id', type=int)
    if not client_id:
        return jsonify({'has_alert': False, 'items': []})
    client = scoped_clients_query().filter_by(id=client_id).first()
    if not client:
        return jsonify({'has_alert': False, 'items': []})
    return jsonify(build_risk_alert_payload(client.cpf))


def require_collector():
    if not current_user.is_authenticated or not current_user.is_collector():
        flash('Acesso permitido apenas para cobrador.', 'danger')
        return False
    return True


def latest_collector_visit(client_id):
    query = filter_query_by_tenant(CollectorVisit.query.filter_by(client_id=client_id), CollectorVisit)
    return query.order_by(CollectorVisit.created_at.desc(), CollectorVisit.id.desc()).first()


def movement_type_label(code):
    labels = {
        'pagamento': 'Pagamento recebido',
        'emprestimo': 'Empréstimo liberado',
        'despesa': 'Despesa',
        'ajuste_manual': 'Ajuste manual',
    }
    return labels.get(code, (code or '-').replace('_', ' ').title())


def _day_bounds(ref_date):
    start = datetime.combine(ref_date, time.min)
    end = datetime.combine(ref_date, time.max)
    return start, end


def compute_cash_closure_report(ref_date):
    ref_date = ref_date or date.today()
    start_dt, end_dt = _day_bounds(ref_date)

    all_before = CashMovement.query.filter(CashMovement.created_at < start_dt).all()
    day_movements = CashMovement.query.filter(CashMovement.created_at >= start_dt, CashMovement.created_at <= end_dt).order_by(CashMovement.created_at.asc(), CashMovement.id.asc()).all()

    opening_balance = quantize_money(sum((Decimal(m.amount or 0) for m in all_before), Decimal('0.00')))
    total_inflows = quantize_money(sum((Decimal(m.amount or 0) for m in day_movements if Decimal(m.amount or 0) > 0), Decimal('0.00')))
    total_outflows = quantize_money(sum((abs(Decimal(m.amount or 0)) for m in day_movements if Decimal(m.amount or 0) < 0), Decimal('0.00')))
    ending_balance = quantize_money(opening_balance + total_inflows - total_outflows)

    payments = Payment.query.filter(Payment.payment_date == ref_date).order_by(Payment.created_at.asc(), Payment.id.asc()).all()
    expenses = Expense.query.filter(Expense.expense_date == ref_date).order_by(Expense.created_at.asc(), Expense.id.asc()).all()
    collector_visits = CollectorVisit.query.filter(CollectorVisit.created_at >= start_dt, CollectorVisit.created_at <= end_dt).order_by(CollectorVisit.created_at.asc(), CollectorVisit.id.asc()).all()

    paid_client_ids = {p.loan.client_id for p in payments if p.loan}
    paid_clients = sorted({p.loan.client.full_name for p in payments if p.loan and p.loan.client}, key=lambda x: x or '')

    overdue_cards = []
    for loan in Loan.query.filter_by(is_active=True).all():
        snap = loan.snapshot(ref_date=ref_date)
        current_due = Decimal(snap.get('current_due_amount') or 0)
        days_overdue = int(snap.get('days_overdue') or 0)
        if current_due > 0 and loan.client_id not in paid_client_ids:
            overdue_cards.append({
                'client_name': loan.client.full_name,
                'loan_id': loan.id,
                'loan_type': loan_cycle_label(loan),
                'days_overdue': days_overdue,
                'amount_due': quantize_money(current_due),
            })
    overdue_cards.sort(key=lambda item: (-item['days_overdue'], item['client_name']))

    collector_summary = {key: 0 for key in VISIT_STATUS_LABELS.keys()}
    for visit in collector_visits:
        collector_summary[visit.status] = collector_summary.get(visit.status, 0) + 1

    movement_groups = []
    for m in day_movements:
        movement_groups.append({
            'label': movement_type_label(m.movement_type),
            'description': m.description,
            'amount': quantize_money(m.amount),
            'created_at': m.created_at,
            'owner_name': m.owner.display_name() if m.owner else '-',
        })

    closure = DailyCashClosure.query.filter_by(closure_date=ref_date).first()

    return {
        'ref_date': ref_date,
        'opening_balance': opening_balance,
        'total_inflows': total_inflows,
        'total_outflows': total_outflows,
        'ending_balance': ending_balance,
        'payments': payments,
        'paid_clients': paid_clients,
        'overdue_cards': overdue_cards,
        'expenses': expenses,
        'collector_visits': collector_visits,
        'collector_summary': collector_summary,
        'movements': movement_groups,
        'closure': closure,
    }


def wrap_text_lines(pdf, text, max_width, font_name='Helvetica', font_size=10):
    words = str(text or '-').split()
    if not words:
        return ['-']
    lines, current = [], words[0]
    for word in words[1:]:
        test = current + ' ' + word
        if stringWidth(test, font_name, font_size) <= max_width:
            current = test
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def draw_wrapped(pdf, text, x, y, max_width, font_name='Helvetica', font_size=10, leading=13, color=colors.HexColor('#24384d')):
    pdf.setFont(font_name, font_size)
    pdf.setFillColor(color)
    for line in wrap_text_lines(pdf, text, max_width, font_name, font_size):
        pdf.drawString(x, y, line)
        y -= leading
    return y


def build_cash_closure_pdf(report):
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    margin = 42
    y = height - 42

    logo_content = brand_logo_bytes()
    if logo_content:
        try:
            img = ImageReader(BytesIO(logo_content))
            iw, ih = img.getSize()
            target_w = 150
            target_h = target_w * (ih / iw)
            pdf.drawImage(img, margin, y - target_h, width=target_w, height=target_h, preserveAspectRatio=True, mask='auto')
            y -= target_h + 16
        except Exception:
            pass

    pdf.setFont('Helvetica-Bold', 16)
    pdf.setFillColor(colors.HexColor('#16385a'))
    pdf.drawString(margin, y, f'Fechamento de caixa - {fmt_date(report["ref_date"])}')
    y -= 24

    pdf.setFillColor(colors.HexColor('#4b647d'))
    pdf.setFont('Helvetica', 10)
    pdf.drawString(margin, y, f'Gerado em {datetime.now().strftime("%d/%m/%Y %H:%M")}')
    if report.get('closure'):
        pdf.drawRightString(width - margin, y, f'Fechado por {report["closure"].closed_by.display_name() if report["closure"].closed_by else "-"}')
    y -= 22

    cards = [
        ('Saldo inicial', fmt_money(report['opening_balance'])),
        ('Entradas', fmt_money(report['total_inflows'])),
        ('Saídas', fmt_money(report['total_outflows'])),
        ('Saldo final', fmt_money(report['ending_balance'])),
    ]
    card_w = (width - (margin * 2) - 18) / 2
    card_h = 44
    for idx, (label, value) in enumerate(cards):
        col = idx % 2
        row = idx // 2
        x = margin + col * (card_w + 18)
        cy = y - row * (card_h + 10)
        pdf.setFillColor(colors.HexColor('#f6f9fd'))
        pdf.roundRect(x, cy - card_h + 4, card_w, card_h, 10, fill=1, stroke=0)
        pdf.setFillColor(colors.HexColor('#6f8398'))
        pdf.setFont('Helvetica-Bold', 9)
        pdf.drawString(x + 10, cy - 10, label)
        pdf.setFillColor(colors.HexColor('#0f2740'))
        pdf.setFont('Helvetica-Bold', 12)
        pdf.drawString(x + 10, cy - 26, value)
    y -= 112

    def ensure_space(required=70):
        nonlocal y
        if y < margin + required:
            pdf.showPage()
            y = height - margin

    def section(title):
        nonlocal y
        ensure_space(45)
        pdf.setFillColor(colors.HexColor('#16385a'))
        pdf.setFont('Helvetica-Bold', 12)
        pdf.drawString(margin, y, title)
        y -= 14
        pdf.setStrokeColor(colors.HexColor('#d9e5f2'))
        pdf.line(margin, y, width - margin, y)
        y -= 14

    section('Resumo do dia')
    summary_lines = [
        f'Clientes que pagaram: {len(report["paid_clients"])}',
        f'Clientes sem pagamento no dia: {len(report["overdue_cards"])}',
        f'Despesas registradas: {len(report["expenses"])}',
        f'Cobranças dos cobradores: {len(report["collector_visits"])}',
    ]
    for line in summary_lines:
        y = draw_wrapped(pdf, line, margin, y, width - margin * 2, font_name='Helvetica', font_size=10) - 4

    section('Quem pagou')
    if report['payments']:
        for payment in report['payments']:
            ensure_space(38)
            client_name = payment.loan.client.full_name if payment.loan and payment.loan.client else '-'
            line = f'{client_name} - {fmt_money(payment.amount_paid)} - {payment.payment_type.replace("_", " ").title()}'
            y = draw_wrapped(pdf, line, margin, y, width - margin * 2) - 3
    else:
        y = draw_wrapped(pdf, 'Nenhum pagamento registrado no dia.', margin, y, width - margin * 2) - 3

    section('Clientes sem pagamento no dia')
    if report['overdue_cards']:
        for item in report['overdue_cards'][:60]:
            ensure_space(38)
            line = f'{item["client_name"]} - contrato #{item["loan_id"]} - {item["loan_type"]} - atraso {item["days_overdue"]} dia(s) - em aberto {fmt_money(item["amount_due"])}'
            y = draw_wrapped(pdf, line, margin, y, width - margin * 2) - 3
    else:
        y = draw_wrapped(pdf, 'Nenhum cliente pendente para a data selecionada.', margin, y, width - margin * 2) - 3

    section('Despesas')
    if report['expenses']:
        for expense in report['expenses']:
            ensure_space(34)
            line = f'{expense.category} - {fmt_money(expense.amount)} - {expense.observation or "Sem observação"}'
            y = draw_wrapped(pdf, line, margin, y, width - margin * 2) - 3
    else:
        y = draw_wrapped(pdf, 'Nenhuma despesa registrada no dia.', margin, y, width - margin * 2) - 3

    section('Movimentações de caixa')
    if report['movements']:
        for mv in report['movements'][:80]:
            ensure_space(34)
            sign = '+' if Decimal(mv['amount']) >= 0 else '-'
            line = f'{mv["label"]} - {mv["description"]} - {sign}{fmt_money(abs(Decimal(mv["amount"]))).replace("R$ ", "R$ ")} - {mv["owner_name"]}'
            y = draw_wrapped(pdf, line, margin, y, width - margin * 2) - 3
    else:
        y = draw_wrapped(pdf, 'Nenhuma movimentação de caixa no dia.', margin, y, width - margin * 2) - 3

    section('Cobranças dos cobradores')
    if report['collector_visits']:
        for visit in report['collector_visits']:
            ensure_space(34)
            status_label = VISIT_STATUS_LABELS.get(visit.status, visit.status)
            line = f'{visit.client.full_name if visit.client else "-"} - {status_label} - {visit.collector.display_name() if visit.collector else "-"}'
            y = draw_wrapped(pdf, line, margin, y, width - margin * 2) - 3
    else:
        y = draw_wrapped(pdf, 'Nenhuma visita de cobrança registrada no dia.', margin, y, width - margin * 2) - 3

    pdf.save()
    buffer.seek(0)
    return buffer


def build_overdue_client_cards():
    today = date.today()
    cards_by_client = {}
    active_loans = filter_query_by_tenant(Loan.query.filter_by(is_active=True), Loan).all()
    for loan in active_loans:
        snap = loan.snapshot(today)
        if snap.get('status_label') != 'Atrasado':
            continue
        current = cards_by_client.get(loan.client_id)
        if current and (current['snap'].get('days_overdue') or 0) >= (snap.get('days_overdue') or 0):
            continue
        visit = latest_collector_visit(loan.client_id)
        cards_by_client[loan.client_id] = {
            'loan': loan,
            'client': loan.client,
            'snap': snap,
            'visit': visit,
        }
    cards = list(cards_by_client.values())
    cards.sort(key=lambda item: item['snap'].get('days_overdue') or 0, reverse=True)
    return cards

def scoped_clients_query():
    if current_user.is_collector():
        return Client.query.filter(sql_text('1=0'))
    return filter_query_by_tenant(Client.query, Client)

def scoped_loans_query():
    if current_user.is_collector():
        return Loan.query.filter(sql_text('1=0'))
    return filter_query_by_tenant(Loan.query, Loan)

def scoped_payments_query():
    if current_user.is_collector():
        return Payment.query.filter(sql_text('1=0'))
    return filter_query_by_tenant(Payment.query, Payment)

def create_excel_backup(path: Path):
    wb = Workbook()
    ws_clients = wb.active
    ws_clients.title = 'clientes'
    ws_clients.append(['ID', 'Nome', 'CPF', 'WhatsApp', 'Cidade', 'Responsável'])
    for c in Client.query.order_by(Client.id).all():
        ws_clients.append([c.id, c.full_name, c.cpf, c.whatsapp, c.city, c.owner.username])

    ws_loans = wb.create_sheet('emprestimos')
    ws_loans.append(['ID', 'Cliente', 'Tipo', 'Valor', 'Juros %', 'Data empréstimo', 'Vencimento', 'Status', 'Ativo', 'Usuário'])
    for l in Loan.query.order_by(Loan.id).all():
        ws_loans.append([l.id, l.client.full_name, l.loan_type, float(l.amount), float(l.interest_percent), fmt_date(l.loan_date), fmt_date(l.due_date), l.status, 'Sim' if l.is_active else 'Não', l.owner_user.username])

    ws_payments = wb.create_sheet('pagamentos')
    ws_payments.append(['ID', 'Empréstimo', 'Cliente', 'Valor pago', 'Tipo', 'Data', 'Usuário'])
    for p in Payment.query.order_by(Payment.id).all():
        ws_payments.append([p.id, p.loan_id, p.loan.client.full_name, float(p.amount_paid), p.payment_type, fmt_date(p.payment_date), p.owner_payment_user.username])

    ws_users = wb.create_sheet('usuarios')
    ws_users.append(['ID', 'Usuário', 'Perfil', 'Ativo'])
    for u in User.query.order_by(User.id).all():
        ws_users.append([u.id, u.username, u.role, 'Sim' if u.is_active_user else 'Não'])

    ws_cash = wb.create_sheet('caixa')
    ws_cash.append(['Saldo atual'])
    ws_cash.append([float(get_cashbox().current_balance or 0)])

    ws_mov = wb.create_sheet('movimentacoes_caixa')
    ws_mov.append(['ID','Tipo','Valor','Descrição','Ref Tipo','Ref ID','Usuário','Criado em'])
    for m in CashMovement.query.order_by(CashMovement.id).all():
        ws_mov.append([m.id, m.movement_type, float(m.amount), m.description, m.reference_type, m.reference_id, m.owner.username if m.owner else '', m.created_at.strftime('%d/%m/%Y %H:%M') if m.created_at else ''])

    ws_exp = wb.create_sheet('despesas')
    ws_exp.append(['ID','Valor','Categoria','Data','Observação','Usuário'])
    for e in Expense.query.order_by(Expense.id).all():
        ws_exp.append([e.id, float(e.amount), e.category, fmt_date(e.expense_date), e.observation, e.owner.username if e.owner else ''])

    wb.save(path)


def cleanup_old_backups(limit=30):
    for folder in [BACKUP_DB_DIR, BACKUP_EXCEL_DIR]:
        files = sorted(folder.glob('*'), key=lambda p: p.stat().st_mtime, reverse=True)
        for old in files[limit:]:
            old.unlink(missing_ok=True)


def perform_backup(prefix='manual'):
    if IS_VERCEL:
        raise RuntimeError('Backups locais não estão disponíveis no Vercel; configure backups no provedor do banco.')
    ensure_dirs()
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    db_backup = BACKUP_DB_DIR / f'{prefix}_db_{timestamp}.sqlite'
    excel_backup = BACKUP_EXCEL_DIR / f'{prefix}_excel_{timestamp}.xlsx'
    if DB_PATH.exists():
        shutil.copy2(DB_PATH, db_backup)
    create_excel_backup(excel_backup)
    cleanup_old_backups(30)


def open_folder(path: Path):
    try:
        if os.name == 'nt':
            os.startfile(path)  # type: ignore[attr-defined]
        elif os.uname().sysname == 'Darwin':
            subprocess.Popen(['open', str(path)])
        else:
            subprocess.Popen(['xdg-open', str(path)])
        return True
    except Exception:
        return False


@app.route('/')
def index():
    logout_user()
    session.clear()
    return redirect(url_for('login'))


@app.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        user = User.query.filter_by(username=username).first()
        if user and user.check_password(password) and user.is_active_user:
            session.permanent = False
            login_user(user)
            if user.tenant_id and not user.is_super_admin():
                session['active_tenant_id'] = user.tenant_id
            flash('Login realizado com sucesso.', 'success')
            return redirect(url_for('dashboard'))
        flash('Usuário ou senha inválidos.', 'danger')
    return render_template('login.html')


@app.route('/instalador', methods=['GET'])
def desktop_installer():
    return render_template(
        'installer.html',
        download_url='https://github.com/mdcredoficial/SISTEMAMD/releases/latest/download/ControlCred-Setup-x64.exe',
    ), 200, {
        'Cache-Control': 'no-store',
        'X-Robots-Tag': 'noindex',
    }




@app.route('/register', methods=['GET', 'POST'])
def register():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        full_name = request.form.get('full_name', '').strip()
        company_name = request.form.get('company_name', '').strip()
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '').strip()
        whatsapp = request.form.get('whatsapp', '').strip()
        if not full_name or not company_name or not username or not password:
            flash('Preencha nome, empresa, usuário e senha para criar seu acesso.', 'danger')
            return render_template('register.html')
        if username.lower() in RESERVED_USERNAMES:
            flash('Este nome de usuário não está disponível. Escolha outro.', 'danger')
            return render_template('register.html')
        if User.query.filter_by(username=username).first():
            flash('Este usuário já está em uso. Escolha outro.', 'danger')
            return render_template('register.html')
        slug = re.sub(r'[^a-z0-9]+', '-', company_name.lower()).strip('-') or 'empresa'
        base_slug = slug
        counter = 2
        while Tenant.query.filter_by(slug=slug).first():
            slug = f"{base_slug}-{counter}"
            counter += 1
        tenant = Tenant(
            company_name=company_name,
            slug=slug,
            subscription_status='trial',
            trial_ends_on=date.today() + relativedelta(days=7),
            is_active=True,
            notes=f'Auto cadastro - WhatsApp: {whatsapp}' if whatsapp else 'Auto cadastro'
        )
        db.session.add(tenant)
        db.session.flush()
        user = User(full_name=full_name, username=username, role='admin', is_active_user=True, tenant_id=tenant.id)
        user.set_password(password)
        db.session.add(user)
        db.session.commit()
        login_user(user)
        session['active_tenant_id'] = tenant.id
        flash('Cadastro realizado com sucesso. Seu teste grátis de 7 dias já está ativo.', 'success')
        return redirect(url_for('dashboard'))
    return render_template('register.html')

@app.route('/logout')
@login_required
def logout():
    logout_user()
    flash('Sessão encerrada.', 'info')
    return redirect(url_for('login'))


@app.route('/dashboard')
@login_required
def dashboard():
    if current_user.is_super_admin():
        tenants = Tenant.query.order_by(Tenant.company_name.asc()).all()
        return render_template('superadmin_dashboard.html', tenants=tenants)
    if current_user.is_collector():
        cards = build_overdue_client_cards()
        return render_template('collector_dashboard.html', cards=cards)
    loans = scoped_loans_query().all()
    clients_count = scoped_clients_query().count()
    active_loans = [l for l in loans if l.is_active]
    total_emprestado = sum(Decimal(l.amount or 0) for l in loans)
    total_receber = sum(l.snapshot()['updated_amount'] for l in active_loans)
    total_recebido = sum(Decimal(p.amount_paid or 0) for p in scoped_payments_query().all())

    today = date.today()
    due_today = [l for l in active_loans if l.snapshot(today)['status_label'] == 'Pagar hoje']
    overdue = [l for l in active_loans if l.snapshot(today)['status_label'] == 'Atrasado']
    recent_payments = scoped_payments_query().order_by(Payment.payment_date.desc(), Payment.id.desc()).limit(8).all()

    month_names = ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez']
    loan_chart = [0] * 12
    payment_chart = [0] * 12
    for loan in loans:
        if loan.loan_date.year == today.year:
            loan_chart[loan.loan_date.month - 1] += float(loan.amount)
    for pay in scoped_payments_query().all():
        if pay.payment_date.year == today.year:
            payment_chart[pay.payment_date.month - 1] += float(pay.amount_paid)

    cashbox = get_cashbox()
    total_despesas = total_expenses_value()
    movements_query = filter_query_by_tenant(CashMovement.query, CashMovement)
    total_entradas = sum(Decimal(m.amount or 0) for m in movements_query.filter(CashMovement.amount > 0).all())
    total_saidas = sum(abs(Decimal(m.amount or 0)) for m in movements_query.filter(CashMovement.amount < 0).all())

    return render_template(
        'dashboard.html',
        clients_count=clients_count,
        active_loans_count=len(active_loans),
        total_emprestado=total_emprestado,
        total_receber=total_receber,
        total_recebido=total_recebido,
        due_today=due_today,
        overdue=overdue,
        recent_payments=recent_payments,
        chart_labels=month_names,
        loan_chart=loan_chart,
        payment_chart=payment_chart,
        saldo_caixa=Decimal(cashbox.current_balance or 0),
        total_despesas=total_despesas,
        total_entradas=total_entradas,
        total_saidas=total_saidas,
    )



@app.route('/plataforma/tenants')
@login_required
def tenants_list():
    if not require_super_admin():
        return redirect(url_for('dashboard'))
    tenants = Tenant.query.order_by(Tenant.company_name.asc()).all()
    return render_template('tenants_list.html', tenants=tenants)

@app.route('/plataforma/tenants/novo', methods=['GET', 'POST'])
@login_required
def tenant_new():
    if not require_super_admin():
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        company_name = request.form.get('company_name', '').strip()
        slug = (request.form.get('slug', '') or company_name.lower().strip()).strip()
        slug = re.sub(r'[^a-z0-9]+', '-', slug.lower()).strip('-')
        admin_name = request.form.get('admin_full_name', '').strip()
        admin_username = request.form.get('admin_username', '').strip()
        admin_password = request.form.get('admin_password', '').strip()
        subscription_status = request.form.get('subscription_status', 'trial')
        trial_days = int(request.form.get('trial_days') or 7)
        if not company_name or not slug or not admin_username or not admin_password:
            flash('Preencha empresa, slug, usuário e senha do administrador.', 'danger')
            return redirect(url_for('tenant_new'))
        if Tenant.query.filter_by(slug=slug).first():
            flash('Slug já está em uso.', 'danger')
            return redirect(url_for('tenant_new'))
        if User.query.filter_by(username=admin_username).first():
            flash('Usuário do administrador já existe.', 'danger')
            return redirect(url_for('tenant_new'))
        tenant = Tenant(company_name=company_name, slug=slug, subscription_status=subscription_status, is_active=True)
        if subscription_status == 'trial':
            tenant.trial_ends_on = date.today() + relativedelta(days=trial_days)
        elif subscription_status in {'active', 'courtesy'}:
            tenant.subscription_ends_on = date.today() + relativedelta(months=1)
            tenant.is_courtesy = subscription_status == 'courtesy'
        db.session.add(tenant)
        db.session.flush()
        admin_user = User(full_name=admin_name or company_name, username=admin_username, role='admin', is_active_user=True, tenant_id=tenant.id)
        admin_user.set_password(admin_password)
        db.session.add(admin_user)
        db.session.commit()
        flash('Empresa criada com sucesso.', 'success')
        return redirect(url_for('tenants_list'))
    return render_template('tenant_form.html', tenant=None)

@app.route('/plataforma/tenants/<int:tenant_id>/editar', methods=['GET', 'POST'])
@login_required
def tenant_edit(tenant_id):
    if not require_super_admin():
        return redirect(url_for('dashboard'))
    tenant = db.session.get(Tenant, tenant_id)
    if not tenant:
        flash('Empresa não encontrada.', 'danger')
        return redirect(url_for('tenants_list'))
    if tenant.slug == 'principal':
        flash('A empresa principal não pode ser excluída ou bloqueada.', 'danger')
        return redirect(url_for('tenants_list'))
    if request.method == 'POST':
        tenant.company_name = request.form.get('company_name', '').strip() or tenant.company_name
        tenant.subscription_status = request.form.get('subscription_status', tenant.subscription_status)
        tenant.is_active = request.form.get('is_active') == '1'
        tenant.is_courtesy = request.form.get('is_courtesy') == '1'
        sub_end = request.form.get('subscription_ends_on', '').strip()
        trial_end = request.form.get('trial_ends_on', '').strip()
        tenant.subscription_ends_on = datetime.strptime(sub_end, '%Y-%m-%d').date() if sub_end else None
        tenant.trial_ends_on = datetime.strptime(trial_end, '%Y-%m-%d').date() if trial_end else None
        logo_file = request.files.get('brand_logo')
        company_name = request.form.get('company_name', '').strip()
        if logo_file and logo_file.filename:
            save_brand_logo(logo_file, tenant=tenant)
        db.session.commit()
        flash('Empresa atualizada com sucesso.', 'success')
        return redirect(url_for('tenants_list'))
    return render_template('tenant_form.html', tenant=tenant)

@app.route('/plataforma/tenants/<int:tenant_id>/selecionar')
@login_required
def tenant_switch(tenant_id):
    if not require_super_admin():
        return redirect(url_for('dashboard'))
    tenant = db.session.get(Tenant, tenant_id)
    if not tenant:
        flash('Cliente não encontrado.', 'danger')
        return redirect(url_for('tenants_list'))
    session['active_tenant_id'] = tenant.id
    flash(f'Visualizando dados de {tenant.display_name()}.', 'success')
    return redirect(url_for('dashboard'))




@app.route('/plataforma/tenants/<int:tenant_id>/bloquear', methods=['POST'])
@login_required
def tenant_toggle_block(tenant_id):
    if not require_super_admin():
        return redirect(url_for('dashboard'))
    tenant = db.session.get(Tenant, tenant_id)
    if not tenant:
        flash('Cliente não encontrado.', 'danger')
        return redirect(url_for('tenants_list'))
    action = request.form.get('action', '').strip()
    if action == 'block':
        tenant.is_active = False
        tenant.subscription_status = 'blocked'
        flash('Cliente SaaS bloqueado com sucesso.', 'success')
    elif action == 'unblock':
        tenant.is_active = True
        if tenant.subscription_status == 'blocked':
            tenant.subscription_status = 'active'
        flash('Cliente SaaS reativado com sucesso.', 'success')
    db.session.commit()
    return redirect(url_for('tenants_list'))


@app.route('/plataforma/tenants/<int:tenant_id>/excluir', methods=['POST'])
@login_required
def tenant_delete(tenant_id):
    if not require_super_admin():
        return redirect(url_for('dashboard'))
    tenant = db.session.get(Tenant, tenant_id)
    if not tenant:
        flash('Cliente não encontrado.', 'danger')
        return redirect(url_for('tenants_list'))
    mode = request.form.get('delete_mode', 'block').strip()
    if mode == 'block':
        tenant.is_active = False
        tenant.subscription_status = 'blocked'
        db.session.commit()
        flash('Cliente SaaS bloqueado. Use essa opção para manter histórico sem excluir dados.', 'success')
        return redirect(url_for('tenants_list'))
    if session.get('active_tenant_id') == tenant.id:
        session.pop('active_tenant_id', None)
    cascade_delete_tenant_data(tenant)
    db.session.commit()
    flash('Cliente SaaS excluído em cascata com todos os dados do CRM.', 'success')
    return redirect(url_for('tenants_list'))


@app.route('/plataforma/tenants/sair')
@login_required
def tenant_switch_clear():
    if not require_super_admin():
        return redirect(url_for('dashboard'))
    session.pop('active_tenant_id', None)
    flash('Visualização do CRM selecionado encerrada.', 'success')
    return redirect(url_for('dashboard'))


@app.route('/calculadora')
@login_required
def calculator_view():
    if current_user.is_super_admin() or current_user.is_collector():
        return redirect(url_for('dashboard'))
    return render_template('calculator.html')

@app.route('/perfil', methods=['GET', 'POST'])
@login_required
def profile():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        logo_file = request.files.get('brand_logo')
        company_name = request.form.get('company_name', '').strip()
        if username:
            existing = User.query.filter(User.username == username, User.id != current_user.id).first()
            if existing:
                flash('Nome de usuário já está em uso.', 'danger')
                return redirect(url_for('profile'))
            current_user.username = username
        if password:
            current_user.set_password(password)
        current_user.charge_message_template = request.form.get('charge_message_template', '').strip()
        tenant = current_tenant()
        if tenant and company_name:
            tenant.company_name = company_name
        if logo_file and logo_file.filename:
            try:
                save_brand_logo(logo_file, tenant=tenant)
            except ValueError as exc:
                flash(str(exc), 'danger')
                return redirect(url_for('profile'))
        db.session.commit()
        flash('Perfil atualizado com sucesso.', 'success')
        return redirect(url_for('profile'))
    return render_template('profile.html', tenant=current_tenant())


@app.route('/usuarios')
@login_required
def users_list():
    if not require_admin():
        return redirect(url_for('dashboard'))
    users = User.query.filter_by(tenant_id=active_tenant_id()).order_by(User.username).all()
    return render_template('users_list.html', users=users)


@app.route('/usuarios/novo', methods=['GET', 'POST'])
@login_required
def user_new():
    if not require_admin():
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        full_name = request.form.get('full_name', '').strip()
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '').strip()
        role = request.form.get('role', 'funcionario')
        if not username or not password:
            flash('Preencha usuário e senha.', 'danger')
            return redirect(url_for('user_new'))
        if User.query.filter_by(username=username).first():
            flash('Usuário já existe.', 'danger')
            return redirect(url_for('user_new'))
        user = User(full_name=full_name, username=username, role=role, is_active_user=True, tenant_id=active_tenant_id())
        user.set_password(password)
        db.session.add(user)
        db.session.commit()
        flash('Usuário cadastrado.', 'success')
        return redirect(url_for('users_list'))
    return render_template('user_form.html', user=None, selected_role=request.args.get('role', 'funcionario'))


@app.route('/usuarios/<int:user_id>/editar', methods=['GET', 'POST'])
@login_required
def user_edit(user_id):
    if not require_admin():
        return redirect(url_for('dashboard'))
    user = db.session.get(User, user_id)
    if not user:
        flash('Usuário não encontrado.', 'danger')
        return redirect(url_for('users_list'))
    if request.method == 'POST':
        full_name = request.form.get('full_name', '').strip()
        username = request.form.get('username', '').strip()
        role = request.form.get('role', 'funcionario')
        password = request.form.get('password', '').strip()
        user.is_active_user = request.form.get('is_active_user') == '1'
        existing = User.query.filter(User.username == username, User.id != user.id).first()
        if existing:
            flash('Usuário já existe.', 'danger')
            return redirect(url_for('user_edit', user_id=user.id))
        user.full_name = full_name
        user.username = username
        user.role = role
        if password:
            user.set_password(password)
        db.session.commit()
        flash('Usuário atualizado.', 'success')
        return redirect(url_for('users_list'))
    return render_template('user_form.html', user=user, selected_role=user.role)


@app.route('/usuarios/<int:user_id>/excluir', methods=['POST'])
@login_required
def user_delete(user_id):
    if not require_admin():
        return redirect(url_for('dashboard'))
    user = db.session.get(User, user_id)
    if not user or user.id == current_user.id:
        flash('Não foi possível excluir este usuário.', 'danger')
        return redirect(url_for('users_list'))
    db.session.delete(user)
    db.session.commit()
    flash('Usuário excluído.', 'success')
    return redirect(url_for('users_list'))




@app.route('/cobradores')
@login_required
def collectors_admin():
    if not require_admin():
        return redirect(url_for('dashboard'))
    collectors = User.query.filter_by(role='cobrador', tenant_id=active_tenant_id()).order_by(User.full_name.asc(), User.username.asc()).all()
    cards = build_overdue_client_cards()
    return render_template('collectors_admin.html', collectors=collectors, cards=cards)


@app.route('/cobrador/clientes/<int:client_id>')
@login_required
def collector_client_detail(client_id):
    if not require_collector():
        return redirect(url_for('dashboard'))
    client = scoped_clients_query().filter_by(id=client_id).first_or_404()
    loans = scoped_loans_query().filter_by(client_id=client.id, is_active=True).order_by(Loan.created_at.desc()).all()
    overdue_loans = []
    for loan in loans:
        snap = loan.snapshot()
        if snap.get('status_label') == 'Atrasado':
            overdue_loans.append((loan, snap))
    if not overdue_loans:
        flash('Este cliente não possui dívida atrasada no momento.', 'warning')
        return redirect(url_for('dashboard'))
    latest_visit = latest_collector_visit(client.id)
    return render_template('collector_client_detail.html', client=client, overdue_loans=overdue_loans, latest_visit=latest_visit)


@app.route('/cobrador/clientes/<int:client_id>/visita', methods=['POST'])
@login_required
def collector_visit_update(client_id):
    if not require_collector():
        return redirect(url_for('dashboard'))
    client = scoped_clients_query().filter_by(id=client_id).first_or_404()
    status = request.form.get('status', '').strip()
    if status not in VISIT_STATUS_LABELS:
        flash('Status de visita inválido.', 'danger')
        return redirect(url_for('collector_client_detail', client_id=client.id))
    visit = CollectorVisit(client_id=client.id, collector_id=current_user.id, status=status, tenant_id=client.tenant_id)
    db.session.add(visit)
    db.session.commit()
    flash('Status da visita registrado com sucesso.', 'success')
    return redirect(url_for('collector_client_detail', client_id=client.id))
@app.route('/clientes')
@login_required
def clients_list():
    query = scoped_clients_query()
    search = request.args.get('q', '').strip()
    if search:
        like = f'%{search}%'
        query = query.filter(
            or_(Client.full_name.ilike(like), Client.cpf.ilike(like), Client.whatsapp.ilike(like), Client.phone.ilike(like))
        )
    clients = query.order_by(Client.full_name).all()
    return render_template('clients_list.html', clients=clients, search=search)


def validate_client_form(form, client=None):
    cpf = format_cpf(form.get('cpf', ''))
    if len(digits_only(cpf)) != 11:
        raise ValueError('CPF inválido.')
    existing = Client.query.filter(Client.cpf == cpf)
    tenant_id = active_tenant_id()
    if tenant_id:
        existing = existing.filter(Client.tenant_id == tenant_id)
    if client:
        existing = existing.filter(Client.id != client.id)
    if existing.first():
        raise ValueError('CPF já cadastrado neste CRM.')

    whatsapp = normalize_whatsapp(form.get('whatsapp', ''))
    if not whatsapp.startswith('+55'):
        raise ValueError('WhatsApp inválido.')

    street = form.get('street', '').strip()
    street_number = form.get('street_number', '').strip()
    neighborhood = form.get('neighborhood', '').strip()
    city = form.get('city', '').strip()
    cep = form.get('cep', '').strip()
    if not all([form.get('full_name', '').strip(), street, street_number, neighborhood, city, cep]):
        raise ValueError('Preencha nome, endereço completo e CPF corretamente.')

    house_status = form.get('house_status', '').strip()
    if house_status and house_status not in {'Própria', 'Aluguel'}:
        raise ValueError('Selecione uma opção válida em Casa própria ou aluguel.')

    return {
        'full_name': form.get('full_name', '').strip(),
        'cpf': cpf,
        'street': street,
        'street_number': street_number,
        'neighborhood': neighborhood,
        'city': city,
        'cep': cep,
        'whatsapp': whatsapp,
        'phone': format_phone(form.get('phone', '')),
        'contact_name': form.get('contact_name', '').strip(),
        'kin_phone': format_phone(form.get('kin_phone', '')),
        'relationship': form.get('relationship', '').strip(),
        'family_street': form.get('family_street', '').strip(),
        'family_number': form.get('family_number', '').strip(),
        'family_neighborhood': form.get('family_neighborhood', '').strip(),
        'family_city': form.get('family_city', '').strip(),
        'family_cep': form.get('family_cep', '').strip(),
        'family_address': form.get('family_address', '').strip(),
        'house_status': house_status,
        'facebook': form.get('facebook', '').strip(),
        'instagram': form.get('instagram', '').strip(),
        'email': form.get('email', '').strip(),
    }


@app.route('/clientes/novo', methods=['GET', 'POST'])
@login_required
def client_new():
    form_data = request.form.to_dict() if request.method == 'POST' else {}
    if request.method == 'POST':
        try:
            data = validate_client_form(request.form)
            alerts = risk_alert_for_cpf(data.get('cpf'))
            if alerts and request.form.get('risk_confirmed') != '1':
                flash('Confirme o alerta de risco para continuar com o cadastro.', 'danger')
                return render_template('client_form.html', client=None, form_data=form_data)
            client = Client(owner_id=current_user.id, tenant_id=active_tenant_id(), **data)
            db.session.add(client)
            db.session.commit()
            for field in DOCUMENT_FIELDS:
                setattr(client, field, save_upload(request.files.get(field), client.id, field))
            db.session.commit()
            flash('Cliente cadastrado com sucesso.', 'success')
            return redirect(url_for('client_detail', client_id=client.id))
        except ValueError as e:
            flash(str(e), 'danger')
    return render_template('client_form.html', client=None, form_data=form_data)


@app.route('/clientes/<int:client_id>')
@login_required
def client_detail(client_id):
    client = scoped_clients_query().filter_by(id=client_id).first_or_404()
    loans = scoped_loans_query().filter_by(client_id=client.id).order_by(Loan.created_at.desc()).all()
    payments = Payment.query.join(Loan).filter(Loan.client_id == client.id)
    if not current_user.is_admin():
        payments = payments.filter(Payment.owner_id == current_user.id)
    payments = payments.order_by(Payment.payment_date.desc(), Payment.id.desc()).all()
    return render_template('client_detail.html', client=client, loans=loans, payments=payments)


@app.route('/clientes/<int:client_id>/editar', methods=['GET', 'POST'])
@login_required
def client_edit(client_id):
    client = scoped_clients_query().filter_by(id=client_id).first_or_404()
    form_data = request.form.to_dict() if request.method == 'POST' else {}
    if request.method == 'POST':
        try:
            data = validate_client_form(request.form, client=client)
            for key, value in data.items():
                setattr(client, key, value)
            db.session.commit()
            for field in DOCUMENT_FIELDS:
                if request.form.get(f'remove_{field}') == '1':
                    remove_client_document(client, field)
                uploaded = request.files.get(field)
                setattr(client, field, save_upload(uploaded, client.id, field, getattr(client, field)))
            db.session.commit()
            flash('Cliente atualizado com sucesso.', 'success')
            return redirect(url_for('client_detail', client_id=client.id))
        except ValueError as e:
            flash(str(e), 'danger')
    return render_template('client_form.html', client=client, form_data=form_data)



@app.route('/clientes/<int:client_id>/documento/<field_name>/remover', methods=['POST'])
@login_required
def client_document_remove(client_id, field_name):
    client = scoped_clients_query().filter_by(id=client_id).first_or_404()
    try:
        remove_client_document(client, field_name)
        db.session.commit()
        flash(f'{DOCUMENT_FIELDS[field_name]} removido com sucesso.', 'success')
    except ValueError as e:
        flash(str(e), 'danger')
    return redirect(request.referrer or url_for('client_detail', client_id=client.id))


@app.route('/clientes/<int:client_id>/excluir', methods=['POST'])
@login_required
def client_delete(client_id):
    if not require_admin():
        return redirect(url_for('clients_list'))
    client = scoped_clients_query().filter_by(id=client_id).first_or_404()
    for field in DOCUMENT_FIELDS:
        delete_document_file(getattr(client, field))
    db.session.delete(client)
    db.session.commit()
    flash('Cliente excluído.', 'success')
    return redirect(url_for('clients_list'))


@app.route('/emprestimos')
@login_required
def loans_list():
    loans = scoped_loans_query().order_by(Loan.created_at.desc()).all()
    return render_template('loans_list.html', loans=loans)


@app.route('/emprestimos/novo', methods=['GET', 'POST'])
@login_required
def loan_new():
    clients = scoped_clients_query().order_by(Client.full_name).all()
    today_str = date.today().isoformat()
    due_default = month_add(date.today()).isoformat()
    form_data = request.form.to_dict() if request.method == 'POST' else {}
    if request.method == 'POST':
        client_id = int(request.form.get('client_id'))
        client = scoped_clients_query().filter_by(id=client_id).first()
        if not client:
            flash('Cliente inválido.', 'danger')
            return redirect(url_for('loan_new'))
        loan_date = datetime.strptime(request.form.get('loan_date'), '%Y-%m-%d').date()
        loan_type = request.form.get('loan_type', 'mensal')
        term_days = int(request.form.get('term_days') or 30)
        term_weeks = int(request.form.get('term_weeks') or 4)
        late_fee_amount = Decimal((request.form.get('late_fee_amount') or '20').replace(',', '.'))
        late_fee_interval_days = int(request.form.get('late_fee_interval_days') or (2 if loan_type == 'diario' else 1))
        charge_on_sundays = request.form.get('charge_on_sundays') == '1'
        due_date = datetime.strptime(request.form.get('due_date'), '%Y-%m-%d').date()
        alerts = risk_alert_for_cpf(client.cpf)
        if alerts and request.form.get('risk_confirmed') != '1':
            flash('Confirme o alerta de risco para continuar com o empréstimo.', 'danger')
            return render_template('loan_form.html', loan=None, clients=clients, today_str=today_str, due_default=due_default, form_data=form_data)
        loan = Loan(
            tenant_id=active_tenant_id(),
            client_id=client.id,
            owner_id=current_user.id,
            amount=Decimal(request.form.get('amount', '0').replace(',', '.')),
            interest_percent=Decimal(request.form.get('interest_percent', '0').replace(',', '.')),
            loan_date=loan_date,
            due_date=due_date,
            loan_type=loan_type,
            term_days=term_days,
            term_weeks=term_weeks,
            late_fee_amount=late_fee_amount,
            late_fee_interval_days=late_fee_interval_days,
            charge_on_sundays=charge_on_sundays,
            status='ativo',
            is_active=True,
            observations=request.form.get('observations', ''),
        )
        db.session.add(loan)
        db.session.flush()
        apply_cash_movement(-Decimal(loan.amount or 0), 'emprestimo', f'Empréstimo #{loan.id} para {loan.client.full_name}', reference_type='loan', reference_id=loan.id)
        sync_shared_risk_for_cpf(client.cpf)
        db.session.commit()
        flash('Empréstimo cadastrado com sucesso.', 'success')
        return redirect(url_for('loans_list'))
    return render_template('loan_form.html', loan=None, clients=clients, today_str=today_str, due_default=due_default, form_data=form_data)


@app.route('/emprestimos/<int:loan_id>/editar', methods=['GET', 'POST'])
@login_required
def loan_edit(loan_id):
    loan = scoped_loans_query().filter_by(id=loan_id).first_or_404()
    clients = scoped_clients_query().order_by(Client.full_name).all()
    if request.method == 'POST':
        client_id = int(request.form.get('client_id'))
        client = scoped_clients_query().filter_by(id=client_id).first()
        if not client:
            flash('Cliente inválido.', 'danger')
            return redirect(url_for('loan_edit', loan_id=loan.id))
        loan.client_id = client.id
        loan.amount = Decimal(request.form.get('amount', '0').replace(',', '.'))
        loan.interest_percent = Decimal(request.form.get('interest_percent', '0').replace(',', '.'))
        loan.loan_date = datetime.strptime(request.form.get('loan_date'), '%Y-%m-%d').date()
        loan.loan_type = request.form.get('loan_type', 'mensal')
        loan.term_days = int(request.form.get('term_days') or 30)
        loan.term_weeks = int(request.form.get('term_weeks') or 4)
        loan.late_fee_amount = Decimal((request.form.get('late_fee_amount') or '20').replace(',', '.'))
        loan.late_fee_interval_days = int(request.form.get('late_fee_interval_days') or (2 if loan.loan_type == 'diario' else 1))
        loan.charge_on_sundays = request.form.get('charge_on_sundays') == '1'
        loan.due_date = datetime.strptime(request.form.get('due_date'), '%Y-%m-%d').date()
        loan.observations = request.form.get('observations', '')
        sync_loan_cash_movement(loan)
        db.session.commit()
        flash('Empréstimo atualizado.', 'success')
        return redirect(url_for('loans_list'))
    return render_template('loan_form.html', loan=loan, clients=clients, today_str=loan.loan_date.isoformat(), due_default=loan.due_date.isoformat(), form_data={})


@app.route('/emprestimos/<int:loan_id>/excluir', methods=['POST'])
@login_required
def loan_delete(loan_id):
    if not require_admin():
        return redirect(url_for('loans_list'))
    loan = scoped_loans_query().filter_by(id=loan_id).first_or_404()
    reverse_cash_movements('loan', loan.id)
    db.session.delete(loan)
    db.session.commit()
    flash('Empréstimo excluído.', 'success')
    return redirect(url_for('loans_list'))


def register_payment(loan, amount_paid, payment_type, payment_date, observation=''):
    amount_paid = Decimal(amount_paid)
    payment = Payment(
        tenant_id=loan.tenant_id,
        loan_id=loan.id,
        owner_id=current_user.id,
        amount_paid=amount_paid,
        payment_type=payment_type,
        payment_date=payment_date,
        observation=observation,
        prev_due_date=loan.due_date,
        prev_status=loan.status,
        prev_is_active=loan.is_active,
        prev_observations=loan.observations,
    )
    if payment_type in ['juros_renovar', 'negociar_juros']:
        # REGRA CONGELADA: a renovação por juros deve obedecer às regras oficiais.
        loan.due_date = get_next_due_date_after_interest_renegotiation(loan, payment_date=payment_date)
        loan.status = 'ativo'
        loan.is_active = True
    elif payment_type in ['quitacao', 'negociar_quitacao']:
        loan.status = 'quitado'
        loan.is_active = False
    else:
        loan.status = 'ativo'
        loan.is_active = True

    db.session.add(payment)
    db.session.flush()
    apply_cash_movement(amount_paid, 'pagamento', f'Pagamento #{payment.id} do contrato #{loan.id}', reference_type='payment', reference_id=payment.id)

    total_paid = quantize_money(sum(Decimal(p.amount_paid or 0) for p in loan.payments))
    total_contract = quantize_money(Decimal(loan.amount or 0) + loan.interest_amount())
    if payment_type in ['parcela_diaria', 'parcela_semanal', 'parcelas_multiplas_diario', 'parcelas_multiplas_semanal'] and total_paid >= total_contract:
        loan.status = 'quitado'
        loan.is_active = False

    db.session.commit()
    return payment


@app.route('/emprestimos/<int:loan_id>/acao/<action>', methods=['POST'])
@login_required
def loan_action(loan_id, action):
    loan = scoped_loans_query().filter_by(id=loan_id).first_or_404()
    snap = loan.snapshot()
    payment_date = datetime.strptime(request.form.get('payment_date', date.today().isoformat()), '%Y-%m-%d').date()
    observation = request.form.get('observation', '')

    if action == 'juros_renovar':
        register_payment(loan, snap['renewal_amount'], 'juros_renovar', payment_date, observation)
        flash('Pagamento de juros registrado e contrato renovado.', 'success')
    elif action == 'quitacao':
        register_payment(loan, snap['settlement_amount'], 'quitacao', payment_date, observation)
        flash('Quitação registrada e contrato encerrado.', 'success')
    elif action == 'negociar_juros':
        negotiated = Decimal(request.form.get('negotiated_amount', '0').replace(',', '.'))
        register_payment(loan, negotiated, 'negociar_juros', payment_date, observation)
        flash('Juros negociados registrados e contrato renovado.', 'success')
    elif action == 'negociar_quitacao':
        negotiated = Decimal(request.form.get('negotiated_amount', '0').replace(',', '.'))
        register_payment(loan, negotiated, 'negociar_quitacao', payment_date, observation)
        flash('Quitação negociada registrada e contrato encerrado.', 'success')
    elif action == 'parcela_diaria':
        amount = Decimal(request.form.get('parcel_amount', '0').replace(',', '.'))
        register_payment(loan, amount, 'parcela_diaria', payment_date, observation)
        flash('Pagamento da parcela diária registrado.', 'success')
    elif action == 'parcela_semanal':
        amount = Decimal(request.form.get('parcel_amount', '0').replace(',', '.'))
        register_payment(loan, amount, 'parcela_semanal', payment_date, observation)
        flash('Pagamento da parcela semanal registrado.', 'success')
    elif action == 'parcelas_multiplas_diario':
        quantity = int(request.form.get('installments_quantity') or 1)
        amount = calculate_multi_installment_amount(loan, quantity, ref_date=payment_date)
        register_payment(loan, amount, 'parcelas_multiplas_diario', payment_date, observation)
        flash(f'Pagamento de {quantity} parcela(s) diária(s) registrado.', 'success')
    elif action == 'parcelas_multiplas_semanal':
        quantity = int(request.form.get('installments_quantity') or 1)
        amount = calculate_multi_installment_amount(loan, quantity, ref_date=payment_date)
        register_payment(loan, amount, 'parcelas_multiplas_semanal', payment_date, observation)
        flash(f'Pagamento de {quantity} parcela(s) semanal(is) registrado.', 'success')
    else:
        flash('Ação inválida.', 'danger')
    next_url = request.form.get('next') or url_for('loans_list')
    return redirect(next_url)


@app.route('/pagamentos')
@login_required
def payments_list():
    payments = scoped_payments_query().order_by(Payment.payment_date.desc(), Payment.id.desc()).all()
    return render_template('payments_list.html', payments=payments)


@app.route('/pagamentos/<int:payment_id>/excluir', methods=['POST'])
@login_required
def payment_delete(payment_id):
    if not require_admin():
        return redirect(url_for('payments_list'))
    payment = scoped_payments_query().filter_by(id=payment_id).first_or_404()
    loan = payment.loan
    loan.due_date = payment.prev_due_date
    loan.status = payment.prev_status
    loan.is_active = payment.prev_is_active
    loan.observations = payment.prev_observations
    reverse_cash_movements('payment', payment.id)
    db.session.delete(payment)
    db.session.commit()
    flash('Pagamento excluído e contrato restaurado ao estado anterior.', 'success')
    return redirect(url_for('payments_list'))




def update_cash_movement_amount(movement, new_amount):
    new_amount = Decimal(new_amount or 0)
    cashbox = get_cashbox()
    current_amount = Decimal(movement.amount or 0)
    delta = new_amount - current_amount
    cashbox.current_balance = quantize_money(Decimal(cashbox.current_balance or 0) + delta)
    movement.amount = quantize_money(new_amount)


def can_manage_cash_movement(movement):
    return movement and movement.movement_type == 'ajuste_manual'

@app.route('/caixa', methods=['GET', 'POST'])
@login_required
def cashbox_view():
    cashbox = get_cashbox()
    report_date_str = request.args.get('report_date') or date.today().isoformat()
    try:
        report_date = datetime.strptime(report_date_str, '%Y-%m-%d').date()
    except ValueError:
        report_date = date.today()

    if request.method == 'POST':
        form_action = request.form.get('form_action', 'adjust')
        if form_action == 'close_day':
            if not current_user.is_admin():
                flash('Apenas administradores podem fechar o caixa.', 'danger')
                return redirect(url_for('cashbox_view', report_date=report_date.isoformat()))
            try:
                closure_date = datetime.strptime(request.form.get('closure_date', report_date.isoformat()), '%Y-%m-%d').date()
            except ValueError:
                closure_date = report_date
            report = compute_cash_closure_report(closure_date)
            closure = DailyCashClosure.query.filter_by(closure_date=closure_date).first()
            if not closure:
                closure = DailyCashClosure(closure_date=closure_date, closed_by_id=current_user.id)
                db.session.add(closure)
            closure.opening_balance = report['opening_balance']
            closure.total_inflows = report['total_inflows']
            closure.total_outflows = report['total_outflows']
            closure.ending_balance = report['ending_balance']
            closure.notes = request.form.get('closure_notes', '').strip() or None
            closure.closed_by_id = current_user.id
            closure.closed_at = datetime.utcnow()
            db.session.commit()
            flash('Fechamento do dia registrado com sucesso.', 'success')
            return redirect(url_for('cashbox_view', report_date=closure_date.isoformat()))

        amount = Decimal(request.form.get('amount', '0').replace(',', '.'))
        movement_kind = request.form.get('movement_kind', 'entrada')
        description = request.form.get('description', '').strip() or 'Ajuste manual de caixa'
        if amount <= 0:
            flash('Informe um valor válido.', 'danger')
            return redirect(url_for('cashbox_view', report_date=report_date.isoformat()))
        signed_amount = amount if movement_kind == 'entrada' else -amount
        apply_cash_movement(signed_amount, 'ajuste_manual', description, reference_type='cash_adjustment')
        db.session.commit()
        flash('Caixa atualizado com sucesso.', 'success')
        return redirect(url_for('cashbox_view', report_date=report_date.isoformat()))
    movements_query = filter_query_by_tenant(CashMovement.query, CashMovement)
    movements = movements_query.order_by(CashMovement.created_at.desc(), CashMovement.id.desc()).limit(80).all()
    closure_report = compute_cash_closure_report(report_date)
    return render_template('cashbox.html', cashbox=cashbox, movements=movements, closure_report=closure_report, report_date=report_date)




@app.route('/caixa/movimentos/<int:movement_id>/editar', methods=['POST'])
@login_required
def cash_movement_edit(movement_id):
    movement = CashMovement.query.get_or_404(movement_id)
    if not can_manage_cash_movement(movement):
        flash('Apenas ajustes manuais podem ser editados.', 'danger')
        return redirect(url_for('cashbox_view'))
    if not current_user.is_admin() and movement.owner_id != current_user.id:
        flash('Você não pode editar esta movimentação.', 'danger')
        return redirect(url_for('cashbox_view'))
    amount = Decimal(request.form.get('amount', '0').replace(',', '.'))
    movement_kind = request.form.get('movement_kind', 'entrada')
    description = request.form.get('description', '').strip() or 'Ajuste manual de caixa'
    if amount <= 0:
        flash('Informe um valor válido.', 'danger')
        return redirect(url_for('cashbox_view'))
    signed_amount = amount if movement_kind == 'entrada' else -amount
    update_cash_movement_amount(movement, signed_amount)
    movement.description = description
    db.session.commit()
    flash('Movimentação atualizada.', 'success')
    return redirect(url_for('cashbox_view'))


@app.route('/caixa/movimentos/<int:movement_id>/excluir', methods=['POST'])
@login_required
def cash_movement_delete(movement_id):
    movement = CashMovement.query.get_or_404(movement_id)
    if not can_manage_cash_movement(movement):
        flash('Apenas ajustes manuais podem ser removidos.', 'danger')
        return redirect(url_for('cashbox_view'))
    if not current_user.is_admin() and movement.owner_id != current_user.id:
        flash('Você não pode excluir esta movimentação.', 'danger')
        return redirect(url_for('cashbox_view'))
    cashbox = get_cashbox()
    cashbox.current_balance = quantize_money(Decimal(cashbox.current_balance or 0) - Decimal(movement.amount or 0))
    db.session.delete(movement)
    db.session.commit()
    flash('Movimentação removida.', 'success')
    return redirect(url_for('cashbox_view'))

@app.route('/caixa/fechamento/pdf')
@login_required
def cash_closure_pdf():
    report_date_str = request.args.get('report_date') or date.today().isoformat()
    try:
        report_date = datetime.strptime(report_date_str, '%Y-%m-%d').date()
    except ValueError:
        report_date = date.today()
    report = compute_cash_closure_report(report_date)
    pdf_buffer = build_cash_closure_pdf(report)
    filename = f"fechamento_caixa_{report_date.strftime('%Y_%m_%d')}.pdf"
    return send_file(pdf_buffer, mimetype='application/pdf', as_attachment=True, download_name=filename)


@app.route('/despesas', methods=['GET', 'POST'])
@login_required
def expenses_list():
    if request.method == 'POST':
        amount = Decimal(request.form.get('amount', '0').replace(',', '.'))
        category = request.form.get('category', '').strip()
        expense_date = datetime.strptime(request.form.get('expense_date', date.today().isoformat()), '%Y-%m-%d').date()
        observation = request.form.get('observation', '').strip()
        if amount <= 0 or not category:
            flash('Informe valor e categoria da despesa.', 'danger')
            return redirect(url_for('expenses_list'))
        expense = Expense(tenant_id=active_tenant_id(), amount=amount, category=category, expense_date=expense_date, observation=observation, owner_id=current_user.id)
        db.session.add(expense)
        db.session.flush()
        apply_cash_movement(-amount, 'despesa', f'Despesa: {category}', reference_type='expense', reference_id=expense.id)
        db.session.commit()
        flash('Despesa registrada com sucesso.', 'success')
        return redirect(url_for('expenses_list'))
    expenses = filter_query_by_tenant(Expense.query, Expense).order_by(Expense.expense_date.desc(), Expense.id.desc()).all()
    return render_template('expenses.html', expenses=expenses, total_despesas=total_expenses_value(), today_str=date.today().isoformat())


@app.route('/despesas/<int:expense_id>/excluir', methods=['POST'])
@login_required
def expense_delete(expense_id):
    if not require_admin():
        return redirect(url_for('expenses_list'))
    expense = filter_query_by_tenant(Expense.query, Expense).filter_by(id=expense_id).first_or_404()
    reverse_cash_movements('expense', expense.id)
    db.session.delete(expense)
    db.session.commit()
    flash('Despesa excluída e caixa restaurado.', 'success')
    return redirect(url_for('expenses_list'))


@app.route('/backups')
@login_required
def backups_list():
    if not require_admin():
        return redirect(url_for('dashboard'))
    if IS_VERCEL:
        return render_template('backups.html', items=[], cloud_backup_mode=True)
    items = []
    for kind, folder in [('db', BACKUP_DB_DIR), ('excel', BACKUP_EXCEL_DIR)]:
        for file in folder.glob('*'):
            stat = file.stat()
            items.append({
                'kind': kind,
                'name': file.name,
                'size': stat.st_size,
                'modified': datetime.fromtimestamp(stat.st_mtime),
            })
    items.sort(key=lambda x: x['modified'], reverse=True)
    return render_template('backups.html', items=items, cloud_backup_mode=False)


@app.route('/backups/gerar', methods=['POST'])
@login_required
def backups_generate():
    if not require_admin():
        return redirect(url_for('dashboard'))
    if IS_VERCEL:
        flash('Os backups online são gerenciados pelo provedor PostgreSQL.', 'warning')
        return redirect(url_for('backups_list'))
    perform_backup('manual')
    flash('Backup gerado com sucesso.', 'success')
    return redirect(url_for('backups_list'))


@app.route('/backups/download/<kind>/<path:filename>')
@login_required
def backup_download(kind, filename):
    if not require_admin():
        return redirect(url_for('dashboard'))
    if IS_VERCEL:
        abort(404)
    folder = BACKUP_DB_DIR if kind == 'db' else BACKUP_EXCEL_DIR
    return send_from_directory(folder, filename, as_attachment=True)


@app.route('/backups/excluir/<kind>/<path:filename>', methods=['POST'])
@login_required
def backup_delete(kind, filename):
    if not require_admin():
        return redirect(url_for('dashboard'))
    if IS_VERCEL:
        abort(404)
    folder = BACKUP_DB_DIR if kind == 'db' else BACKUP_EXCEL_DIR
    (folder / filename).unlink(missing_ok=True)
    flash('Backup excluído.', 'success')
    return redirect(url_for('backups_list'))


@app.route('/backups/abrir/<kind>')
@login_required
def backup_open_folder(kind):
    if not require_admin():
        return redirect(url_for('dashboard'))
    if IS_VERCEL:
        flash('Pastas locais não existem no Vercel. Consulte os backups do banco PostgreSQL.', 'warning')
        return redirect(url_for('backups_list'))
    folder = BACKUP_DB_DIR if kind == 'db' else BACKUP_EXCEL_DIR
    if open_folder(folder):
        flash(f'Pasta aberta: {folder}', 'success')
    else:
        flash(f'Não foi possível abrir automaticamente. Caminho: {folder}', 'warning')
    return redirect(url_for('backups_list'))


@app.route('/uploads/<path:filename>')
def uploaded_file(filename):
    key = storage_key(filename)
    parts = key.split('/')
    is_global_brand_asset = parts[:2] == ['brand', 'global']
    if not is_global_brand_asset and not current_user.is_authenticated:
        return redirect(url_for('login'))
    if parts[0] == 'tenants' and len(parts) > 1:
        try:
            file_tenant_id = int(parts[1])
        except ValueError:
            abort(404)
        active_id = session.get('active_tenant_id') if current_user.is_super_admin() else getattr(current_user, 'tenant_id', None)
        if active_id and int(active_id) != file_tenant_id:
            abort(404)
        if not current_user.is_super_admin() and not active_id:
            abort(404)
        if not db.session.get(Tenant, file_tenant_id):
            abort(404)
    elif is_global_brand_asset:
        pass
    else:
        attachment_filter = or_(*(getattr(Client, field) == key for field in DOCUMENT_FIELDS))
        if not filter_query_by_tenant(Client.query.filter(attachment_filter), Client).first():
            abort(404)
    try:
        content = storage_get(key)
    except FileNotFoundError:
        abort(404)
    return send_file(
        BytesIO(content),
        mimetype=mimetypes.guess_type(key)[0] or 'application/octet-stream',
        download_name=Path(key).name,
        conditional=True,
    )


AUTO_INIT_DB = os.getenv('AUTO_INIT_DB', '0' if IS_VERCEL else '1').lower() in {'1', 'true', 'yes'}
if AUTO_INIT_DB:
    with app.app_context():
        ensure_dirs()
        create_initial_data()
        run_migrations()
        if not IS_VERCEL:
            try:
                perform_backup('startup')
            except Exception:
                pass


if not IS_VERCEL:
    import atexit

    @atexit.register
    def shutdown_backup():
        try:
            with app.app_context():
                perform_backup('shutdown')
        except Exception:
            pass


if __name__ == '__main__':
    app.run(debug=True)
